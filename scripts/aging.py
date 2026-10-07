"""Balance âgée clients, DSO et liste de relance à partir des postes ouverts (ex. SAP FBL5N) et, si fourni, du chiffre d'affaires mensuel.
Entrées : --open  postes ouverts : Client, N° document, Type (facture / DN / CN / acompte…), Date document, Échéance, Montant (signé), Devise [, Référence, Texte]
          --sales chiffre d'affaires mensuel (pour le DSO) : Mois (AAAA-MM), Montant [, Client]     (optionnel ; sinon DSO calculé sur les factures des postes ouverts)
Sortie : classeur « Balance âgée » : par client × tranche, par document, DSO (méthode count-back), top débiteurs, liste de relance, synthèse en formules.
Usage : aging.py --open FBL5N.xlsx [--sales ca.xlsx] --out balance.xlsx [--arrete AAAA-MM-JJ] [--config alias.json] [--relance 30]"""
import argparse, json
import pandas as pd, numpy as np
from openpyxl import Workbook
from xlsx_common import *

ap = argparse.ArgumentParser(); ap.add_argument("--open", required=True); ap.add_argument("--sales"); ap.add_argument("--out", required=True); ap.add_argument("--arrete"); ap.add_argument("--config"); ap.add_argument("--relance", type=int, default=30, help="retard (j) à partir duquel un document entre dans la liste de relance")
a = ap.parse_args()
DEF = {"open": {"Client": ["CUSTOMER", "NAME 1", "CLIENT NAME", "NOM CLIENT", "ACCOUNT", "SOLD TO"], "N° document": ["DOCUMENT NUMBER", "DOCUMENT", "REFERENCE", "BILLING DOCUMENT", "N° FACTURE", "NUMÉRO FACTURE"], "Type": ["DOCUMENT TYPE", "TYPE", "DOC. TYPE", "TYPE COMMANDE DESCP"],
                "Date document": ["DOCUMENT DATE", "POSTING DATE", "DATE", "DATE CRÉATION FACTURE", "BASELINE DATE"], "Échéance": ["NET DUE DATE", "DUE DATE", "ECHEANCE", "ÉCHÉANCE"], "Montant": ["AMOUNT IN LOCAL CURRENCY", "AMOUNT", "OPEN AMOUNT", "MONTANT", "SOLDE CLIENT", "AMOUNT IN DOC. CURR."],
                "Devise": ["CURRENCY", "DEVISE", "DOCUMENT CURRENCY"], "Référence": ["REFERENCE", "ASSIGNMENT", "RÉFÉRENCE"], "Texte": ["TEXT", "ITEM TEXT", "LIBELLÉ"], "Navire": ["VESSEL", "NAVIRE"]},
       "sales": {"Mois": ["MONTH", "MOIS", "PERIOD", "PÉRIODE"], "Montant": ["AMOUNT", "MONTANT", "NET VALUE", "CA", "REVENUE"], "Client": ["CUSTOMER", "CLIENT"]}}
cfg = load_config(a.config, DEF)
Op = read_any(a.open, aliases=cfg["open"], required=["Client", "Montant"])
ARRETE = pd.Timestamp(a.arrete) if a.arrete else pd.Timestamp.today().normalize()
Op["Montant"] = Op["Montant"].map(to_num); Op = Op[Op["Montant"].notna() & (Op["Montant"].abs() > 0.005)].copy()
for c in ["Date document", "Échéance"]:
    if c in Op.columns: Op[c] = Op[c].map(to_date)
if "Échéance" not in Op.columns: Op["Échéance"] = Op["Date document"] if "Date document" in Op.columns else pd.NaT
if "Date document" not in Op.columns: Op["Date document"] = Op["Échéance"]
Op["Client"] = Op["Client"].astype(str).str.strip().str.upper().str.replace(r"\s+", " ", regex=True)
Op["Retard (j)"] = (ARRETE - Op["Échéance"]).dt.days
Op["Tranche"] = Op["Retard (j)"].map(bucket_days)
Op["Âge document (j)"] = (ARRETE - Op["Date document"]).dt.days
Op["Sens"] = np.where(Op["Montant"] > 0, "Débit (créance)", "Crédit (avoir / acompte)")
Op["Relance"] = np.where((Op["Montant"] > 0) & (Op["Retard (j)"] > a.relance), f"À relancer (> {a.relance} j)", "")
# ---- balance par client × tranche (créances nettes des crédits : on affiche brut débit, crédit, net)
piv = Op[Op["Montant"] > 0].pivot_table(index="Client", columns="Tranche", values="Montant", aggfunc="sum", fill_value=0).reindex(columns=BUCKETS, fill_value=0)
piv["Total débit"] = piv.sum(axis=1)
cred = Op[Op["Montant"] < 0].groupby("Client")["Montant"].sum()
piv["Crédits non imputés"] = cred.reindex(piv.index).fillna(0); piv["Net"] = piv["Total débit"] + piv["Crédits non imputés"]
piv["% > 90 j"] = ((piv["91-180 j"] + piv["> 180 j"]) / piv["Total débit"].replace(0, np.nan)).fillna(0)
piv = piv.sort_values("Net", ascending=False).reset_index()
# ---- DSO count-back : jours de CA couverts par l'encours net
dso = None; dso_detail = []
if a.sales:
    S = read_any(a.sales, aliases=cfg["sales"], required=["Mois", "Montant"]); S["Montant"] = S["Montant"].map(to_num); S["Mois"] = S["Mois"].astype(str).str[:7]
    ca = S.groupby("Mois")["Montant"].sum().sort_index(ascending=False)
elif "Date document" in Op.columns and Op["Date document"].notna().any():
    inv = Op[Op["Montant"] > 0]; ca = inv.groupby(inv["Date document"].dt.to_period("M").astype(str))["Montant"].sum().sort_index(ascending=False)  # approximation : factures ouvertes seulement
else: ca = pd.Series(dtype=float)
enc = float(Op["Montant"].sum()); rest = enc; days = 0.0
for mois, m_ca in ca.items():
    if rest <= 0 or m_ca <= 0: break
    dim = pd.Period(mois).days_in_month
    if rest >= m_ca: days += dim; rest -= m_ca; dso_detail.append((mois, m_ca, dim))
    else: days += dim * rest / m_ca; dso_detail.append((mois, m_ca, round(dim * rest / m_ca, 1))); rest = 0
dso = round(days, 1) if len(ca) else None
# ---- classeur
wb = Workbook(); ws = wb.active; ws.title = "Synthèse"
ws["A1"] = "Balance âgée clients, DSO et relances"; ws["A1"].font = TITLE
ws["A2"] = f"Arrêté : {ARRETE:%d/%m/%Y}. Retard = arrêté − échéance. Tranches sur les documents débiteurs ; crédits non imputés affichés à part. DSO count-back {'sur le CA mensuel fourni' if a.sales else 'approximé sur les factures ouvertes (fournir --sales pour un DSO exact)'}."; ws["A2"].font = NOTE
cm = colmap(Op); n = len(Op); R_ = lambda c: rng("Postes ouverts", cm[c], n)
r = kpi_rows(ws, 4, [("Postes ouverts", f'=COUNTA({R_("Client")})', "0"), ("Encours débiteur (créances)", f'=SUMIF({R_("Montant")},">0")', MONEY0), ("Crédits non imputés", f'=SUMIF({R_("Montant")},"<0")', MONEY0), ("Encours net", f'=SUM({R_("Montant")})', MONEY0),
                    ("Clients débiteurs", len(piv), "0"), ("DSO count-back (jours)", dso if dso is not None else "n/a", "0.0")], "1. Encours")
r = kpi_rows(ws, r, [(b, f'=SUMIFS({R_("Montant")},{R_("Tranche")},"{b}",{R_("Montant")},">0")', MONEY0) for b in BUCKETS] + [("Part > 90 j", f'=IF(B{r + 1}+B{r + 2}+B{r + 3}+B{r + 4}+B{r + 5}+B{r + 6}=0,0,(B{r + 5}+B{r + 6})/(B{r + 1}+B{r + 2}+B{r + 3}+B{r + 4}+B{r + 5}+B{r + 6}))', PCT)], "2. Encours vieilli (créances)")
r = kpi_rows(ws, r, [("Documents à relancer", f'=COUNTIF({R_("Relance")},"À relancer*")', "0"), ("Montant à relancer", f'=SUMIFS({R_("Montant")},{R_("Relance")},"À relancer*")', MONEY0)], f"3. Relances (retard > {a.relance} j)")
ws.column_dimensions["A"].width = 40; ws.column_dimensions["B"].width = 18; ws.sheet_view.showGridLines = False
write_df(wb.create_sheet("Par client"), piv, money0=BUCKETS + ["Total débit", "Crédits non imputés", "Net"], pct=["% > 90 j"], widths={"Client": 40})
rel = Op[Op["Relance"] != ""].sort_values(["Client", "Retard (j)"], ascending=[True, False])
cols = [c for c in ["Client", "N° document", "Type", "Référence", "Navire", "Date document", "Échéance", "Montant", "Devise", "Retard (j)", "Tranche", "Texte"] if c in rel.columns]
write_df(wb.create_sheet("Liste de relance"), rel[cols], money=["Montant"], date=["Date document", "Échéance"], intc=["Retard (j)"], widths={"Client": 36, "Texte": 40})
cols = [c for c in ["Client", "N° document", "Type", "Sens", "Référence", "Navire", "Date document", "Échéance", "Montant", "Devise", "Âge document (j)", "Retard (j)", "Tranche", "Relance", "Texte"] if c in Op.columns]
ws4 = wb.create_sheet("Postes ouverts"); write_df(ws4, Op[cols], money=["Montant"], date=["Date document", "Échéance"], intc=["Âge document (j)", "Retard (j)"], widths={"Client": 36, "Texte": 40})
Op = Op[cols]
if dso_detail: write_df(wb.create_sheet("DSO"), pd.DataFrame(dso_detail, columns=["Mois", "CA du mois", "Jours retenus"]), money0=["CA du mois"], autofilter=False)
summary = {"module": "aging", "arrete": str(ARRETE.date()), "postes": n, "encours_debit": float(Op["Montant"][Op["Montant"] > 0].sum()), "credits": float(Op["Montant"][Op["Montant"] < 0].sum()), "encours_net": enc, "dso": dso,
           "tranches": {b: float(Op.loc[(Op["Tranche"] == b) & (Op["Montant"] > 0), "Montant"].sum()) for b in BUCKETS}, "relances": int((Op["Relance"] != "").sum()), "montant_relance": float(Op.loc[Op["Relance"] != "", "Montant"].sum()),
           "top_clients": piv.head(10)[["Client", "Total débit", "Net", "% > 90 j"]].to_dict("records")}
finish(wb, a.out, summary)
print(json.dumps({k: summary[k] for k in ["postes", "encours_debit", "encours_net", "dso", "relances"]}, ensure_ascii=False, default=str))
