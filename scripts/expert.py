"""Analyse experte DEM/DESP : typologie de fuite par voyage, délais, encours vieilli, référentiel navires,
doublons SAP, contrôles automatiques, échantillon d'audit, plan d'action chiffré.
Entrées : classeur de rapprochement (produit par reconcile.py) + export SAP d'origine (pour IMO) + optionnel factures de fret.
Sortie : classeur 'Analyse experte'."""
import re, sys, json, argparse
import pandas as pd
import numpy as np
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule

ap = argparse.ArgumentParser()
ap.add_argument("--rapprochement", required=True)
ap.add_argument("--sap", required=True, help="export SAP d'origine (Feuil1)")
ap.add_argument("--fret", default=None, help="optionnel : factures de fret armateur (modèle fourni)")
ap.add_argument("--out", required=True)
ap.add_argument("--delai-alerte", type=int, default=180, help="jours depuis BL sans DN émise → alerte prescription")
ap.add_argument("--tolerance", type=float, default=1.0)
ap.add_argument("--sap-sheet", default="Feuil1")
ap.add_argument("--sens-config", default=None, help="JSON {type de réclamation: incoterm|incoterm_inverse|creance|dette|signe}")
args = ap.parse_args()

P = args.rapprochement
rap = pd.read_excel(P, sheet_name="Rapprochement")
det = pd.read_excel(P, sheet_name="Détail tracker")
sap = pd.read_excel(P, sheet_name="Base SAP 2024+")
nts = pd.read_excel(P, sheet_name="SAP non suivi")
_xs = pd.ExcelFile(args.sap)
exp = pd.read_excel(args.sap, sheet_name=args.sap_sheet if args.sap_sheet in _xs.sheet_names else _xs.sheet_names[0])
_norm = {re.sub(r"\s+", " ", str(c)).strip().upper(): c for c in exp.columns}
for canon, aliases in {"Numéro Facture": ["BILLING DOCUMENT", "N° FACTURE", "INVOICE"], "Navire": ["VESSEL", "SHIP"], "Navir IMO": ["IMO"], "Report Run Date": ["DATE EXTRACTION"]}.items():
    if canon not in exp.columns:
        for a in aliases:
            if a in _norm: exp = exp.rename(columns={_norm[a]: canon}); break
if "Report Run Date" not in exp.columns: exp["Report Run Date"] = pd.Timestamp.today().normalize()
if "Navir IMO" not in exp.columns: exp["Navir IMO"] = np.nan
sap["Numéro Facture"] = sap["Numéro Facture"].astype(str)
exp["Numéro Facture"] = exp["Numéro Facture"].astype(str)
ARRETE = pd.to_datetime(exp["Report Run Date"].max())  # date d'arrêté = date d'extraction SAP
M = "Montant tracker (FINAL AMOUNT)"; S = "SAP montant refacturé (abs)"; E = "SAP encaissé / imputé (abs)"; SO = "SAP solde client (abs)"
KEY = "Clé (navire | BL | type SAP)"


def norm(v):
    if pd.isna(v): return ""
    s = str(v).upper(); s = re.sub(r"\bM/?V\b", " ", s); s = re.sub(r"[^A-Z0-9 ]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def pay_cat(s):
    s = str(s).upper()
    if "NOT YET" in s: return "NOT YET STARTED"
    if "PAID" in s: return "PAID"
    if "PENDING" in s: return "PENDING"
    return "AUTRE"


# ------------------------------------------------------------------ 0. factures de fret (optionnel)
fret = None
if args.fret:
    f = pd.read_excel(args.fret)
    f["VK"] = f["Navire"].map(norm); f["BL"] = pd.to_datetime(f["Date BL"]).dt.normalize()
    fret = f

# ------------------------------------------------------------------ 1. documents SAP → clé
doc2key = {}
for _, r in rap.iterrows():
    for n in str(r["Docs SAP (n° facture)"]).split(" ; "):
        n = n.strip()
        if n and n != "nan": doc2key[n] = r[KEY]
sap["Clé"] = sap["Numéro Facture"].map(doc2key)
sap["BL"] = pd.to_datetime(sap["Date BL"]); sap["Création"] = pd.to_datetime(sap["Date Création Facture"])
sap["Echéance"] = pd.to_datetime(sap["Echéance"]); sap["Règlement 1"] = pd.to_datetime(sap["Date Reglement 1"])
sap["Délai BL→DN (j)"] = (sap["Création"] - sap["BL"]).dt.days
sap["Délai DN→règlement (j)"] = (sap["Règlement 1"] - sap["Création"]).dt.days
sap["Retard vs échéance (j)"] = np.where(sap["Solde abs"] > 0.01, (ARRETE - sap["Echéance"]).dt.days, np.nan)


def bucket(d):
    if pd.isna(d): return ""
    if d <= 0: return "Non échu"
    if d <= 30: return "1-30 j"
    if d <= 60: return "31-60 j"
    if d <= 90: return "61-90 j"
    if d <= 180: return "91-180 j"
    return "> 180 j"


sap["Tranche retard"] = sap["Retard vs échéance (j)"].map(bucket)

# ------------------------------------------------------------------ 2. chaîne par voyage + typologie
rap["Pay"] = rap["Statut paiement (tracker)"].map(pay_cat)
rap["BL"] = pd.to_datetime(rap["Date BL (début)"])
rap["Âge BL à l'arrêté (j)"] = (ARRETE - rap["BL"]).dt.days
agg = sap[sap["Clé"].notna()].groupby("Clé").agg(
    DN_premiere=("Création", "min"), DN_derniere=("Création", "max"), Regl_premier=("Règlement 1", "min"),
    Echeance_min=("Echéance", "min"), Retard_max=("Retard vs échéance (j)", "max"), Delai_BL_DN=("Délai BL→DN (j)", "min"),
    Delai_DN_regl=("Délai DN→règlement (j)", "min"))
rap = rap.merge(agg, left_on=KEY, right_index=True, how="left")
if fret is not None:
    fa = fret.groupby(["VK", "BL"]).agg(Fret_paye=("Montant demurrage payé", "sum"), Fret_date=("Date paiement", "max")).reset_index()
    rap["VK"] = rap["Navire"].map(norm)
    rap = rap.merge(fa, left_on=["VK", "BL"], right_on=["VK", "BL"], how="left")
else:
    rap["Fret_paye"] = np.nan; rap["Fret_date"] = pd.NaT


def typo(r):
    if r["Nb docs SAP"] == 0:
        if r["Pay"] == "NOT YET STARTED":
            return "T1 – Non réclamé / non démarré" if r["Âge BL à l'arrêté (j)"] <= args.delai_alerte else "T1b – Non démarré > délai d'alerte"
        return "T2 – Réclamé, DN/CN absente de l'export SAP"
    if r["Contrôle écart"] in ("SAP < tracker", "SAP > tracker"):
        return "T3 – Écart de montant tracker / SAP"
    if r[SO] > args.tolerance and r[E] <= args.tolerance:
        return "T4 – Refacturé, non encaissé"
    if r[SO] > args.tolerance:
        return "T5 – Refacturé, partiellement encaissé"
    return "T0 – Soldé"


rap["Typologie"] = rap.apply(typo, axis=1)
rap["Montant exposé (USD)"] = np.where(rap["Nb docs SAP"] == 0, rap[M], rap[SO])
# Sens du flux : en CFR/CIF/DPU le demurrage est une créance (débit note) et le despatch une dette (crédit note) ;
# en FOB c'est l'inverse : le demurrage est dû par le vendeur à l'acheteur (paiement OP / crédit note) et le despatch est une débit note à l'acheteur.
rap["Incoterm"] = rap["Mode livraison"].astype(str).str.upper().str.extract(r"(FOB|CFR|CIF|DPU|DAP|CPT)")[0].fillna("?") if "Mode livraison" in rap.columns else "?"
SENS_RULES = {"DEMURRAGE": "incoterm", "DESPATCH": "incoterm_inverse", "DISPATCH": "incoterm_inverse", "DISCOUNT": "dette", "REMISE": "dette", "PRICE ADJ": "signe", "AJUSTEMENT": "signe",
              "CLAIM": "dette", "LC CHARGE": "creance", "MARGIN SHARE": "dette", "DEAD FREIGHT": "creance", "TROP PERCU": "dette", "VIREMENT": "dette"}
if args.sens_config: SENS_RULES.update({k.upper(): v for k, v in json.load(open(args.sens_config)).items()})
CRE, DET = "Créance (débit note à l'acheteur)", "Dette (crédit note / paiement à l'acheteur)"


def sens(r):
    ct = str(r.get("Type réclamation", "")).upper(); t = r["Type SAP attendu"]
    if t == "ZDS/ZDP" or ("/" in ct and "DEM" in ct and "DESP" in ct): return "Mixte"
    rule = next((v for k, v in SENS_RULES.items() if k in ct), None)
    if rule is None: rule = "incoterm" if t == "ZDS" else ("incoterm_inverse" if t == "ZDP" else "signe")
    if rule == "incoterm": return (DET if t != "ZDP" else CRE) if r["Incoterm"] == "FOB" else (CRE if t != "ZDP" else DET)
    if rule == "incoterm_inverse": return (CRE if r["Incoterm"] == "FOB" else DET)
    if rule == "creance": return CRE
    if rule == "dette": return DET
    # signe : d'après le signe des documents SAP rattachés ; sinon indéterminé
    deb, crd = float(r.get("SAP ZDS – Débit (demurrage)", 0) or 0), float(r.get("SAP ZDP – Crédit (despatch)", 0) or 0)
    if deb > 0 and crd == 0: return CRE
    if crd < 0 and deb == 0: return DET
    return "Indéterminé (à qualifier)"
rap["Sens du flux"] = rap.apply(sens, axis=1)
rap["Créance exposée (USD)"] = np.where(rap["Sens du flux"].str.startswith("Créance"), rap["Montant exposé (USD)"], 0.0)
rap["Dette exposée (USD)"] = np.where(rap["Sens du flux"].str.startswith("Dette"), rap["Montant exposé (USD)"], 0.0)
rap["Alerte doublons SAP"] = np.where(rap["Alerte"].notna(), "Oui", "")
rap["Alerte délai"] = np.where((rap["Nb docs SAP"] == 0) & (rap["Âge BL à l'arrêté (j)"] > args.delai_alerte), f"> {args.delai_alerte} j sans DN", "")
rap["Tranche retard (max)"] = rap["Retard_max"].map(bucket)

chain_cols = {
    KEY: "Clé", "Navire": "Navire", "BL": "Date BL", "Année BL": "Année", "Type SAP attendu": "Type", "Région (tracker)": "Région", "Client (tracker)": "Client",
    "Fret_paye": "① Payé armateur (USD)", "Fret_date": "① Date paiement fret",
    M: "② Réclamé tracker (USD)", "Pay": "② Statut tracker", "DN/CN n° (tracker)": "② DN/CN citées",
    S: "③ Refacturé SAP (USD)", "Nb docs SAP": "③ Nb docs SAP", "DN_premiere": "③ 1ère DN émise", "Delai_BL_DN": "③ Délai BL→DN (j)",
    E: "④ Encaissé / imputé (USD)", SO: "④ Solde client (USD)", "Regl_premier": "④ 1er règlement", "Delai_DN_regl": "④ Délai DN→règlement (j)",
    "Echeance_min": "④ Échéance", "Retard_max": "④ Retard vs échéance (j)", "Tranche retard (max)": "④ Tranche retard",
    "Écart tracker – SAP": "Écart ②–③ (USD)", "Âge BL à l'arrêté (j)": "Âge BL à l'arrêté (j)", "Typologie": "Typologie de fuite",
    "Montant exposé (USD)": "Montant exposé (USD)", "Incoterm": "Incoterm", "Sens du flux": "Sens du flux", "Créance exposée (USD)": "Créance exposée (USD)", "Dette exposée (USD)": "Dette exposée (USD)", "Alerte doublons SAP": "Alerte doublons SAP", "Alerte délai": "Alerte délai", "Méthode de rapprochement": "Méthode",
}
chain = rap[list(chain_cols)].rename(columns=chain_cols).sort_values(["Typologie de fuite", "Montant exposé (USD)"], ascending=[True, False]).reset_index(drop=True)

# ------------------------------------------------------------------ 3. synthèse typologie + scénarios
typ = chain.groupby("Typologie de fuite").agg(Clés=("Clé", "size"), Réclamé=("② Réclamé tracker (USD)", "sum"), Refacturé=("③ Refacturé SAP (USD)", "sum"),
                                              Encaissé=("④ Encaissé / imputé (USD)", "sum"), Exposé=("Montant exposé (USD)", "sum"), Créance=("Créance exposée (USD)", "sum"), Dette=("Dette exposée (USD)", "sum")).reset_index()
T = dict(zip(typ["Typologie de fuite"], typ["Créance"]))   # scénarios de récupération = créances seulement
TD = dict(zip(typ["Typologie de fuite"], typ["Dette"]))
sens_tab = chain.groupby(["Incoterm", "Sens du flux"]).agg(Clés=("Clé", "size"), Réclamé=("② Réclamé tracker (USD)", "sum"), Refacturé=("③ Refacturé SAP (USD)", "sum"), Encaissé=("④ Encaissé / imputé (USD)", "sum"), Exposé=("Montant exposé (USD)", "sum")).reset_index()
scen = pd.DataFrame([
    ["Récupérable certain", "T4 + T5 : refacturé, solde client ouvert", T.get("T4 – Refacturé, non encaissé", 0) + T.get("T5 – Refacturé, partiellement encaissé", 0), "Relance clients / lettrage ; vérifier annulations (doublons)"],
    ["Récupérable probable", "T3 : écarts de montant (part SAP < tracker)", float(chain.loc[(chain["Typologie de fuite"].str.startswith("T3")) & (chain["Écart ②–③ (USD)"] > 0) & chain["Sens du flux"].str.startswith("Créance"), "Écart ②–③ (USD)"].sum()), "Compléter la refacturation ou corriger le tracker"],
    ["À confirmer (export filtré ?)", "T2 : réclamé, DN/CN absente de l'export", T.get("T2 – Réclamé, DN/CN absente de l'export SAP", 0), "Relancer l'export SAP complet ; si confirmé absent → refacturer"],
    ["En cours de traitement", "T1 : non démarré dans le délai", T.get("T1 – Non réclamé / non démarré", 0), "Suivi du délai d'émission"],
    ["Risque de prescription", f"T1b : non démarré > {args.delai_alerte} j", T.get("T1b – Non démarré > délai d'alerte", 0), "Émettre en priorité ; vérifier clause contractuelle de délai"],
    ["Dettes sans document dans l'export", "T2 côté dette : réglées hors débit note (FOB demurrage payé par OP / swift) ou crédit notes absentes de l'export", TD.get("T2 – Réclamé, DN/CN absente de l'export SAP", 0), "Rien à récupérer ; vérifier la comptabilisation et l'exhaustivité de l'export"],
    ["Dettes à régler ou provisionner", "T1 + T1b + T4 + T5 côté dette : crédit notes / paiements dus à l'acheteur non émis ou non imputés", sum(TD.get(k, 0) for k in ["T1 – Non réclamé / non démarré", "T1b – Non démarré > délai d'alerte", "T4 – Refacturé, non encaissé", "T5 – Refacturé, partiellement encaissé"]), "Provision ; émission des crédit notes ; pas une recette"],
], columns=["Scénario", "Définition", "Montant (USD)", "Action"])
scen_sens = "Les scénarios de récupération ne comptent que les créances (débit notes à l'acheteur). Les dettes (crédit notes, paiements FOB) sont isolées : elles ne se récupèrent pas, elles se règlent."

# ------------------------------------------------------------------ 4. délais
d_dn = sap[sap["Clé"].notna()]
delais = pd.DataFrame([
    ["Délai BL → émission DN/CN (jours)", d_dn["Délai BL→DN (j)"].median(), d_dn["Délai BL→DN (j)"].mean(), d_dn["Délai BL→DN (j)"].quantile(0.9), len(d_dn)],
    ["Délai DN/CN → 1er règlement (jours)", d_dn["Délai DN→règlement (j)"].median(), d_dn["Délai DN→règlement (j)"].mean(), d_dn["Délai DN→règlement (j)"].quantile(0.9), int(d_dn["Délai DN→règlement (j)"].notna().sum())],
    ["Âge des dossiers non démarrés (jours depuis BL)", rap.loc[rap["Typologie"].str.startswith("T1"), "Âge BL à l'arrêté (j)"].median(), rap.loc[rap["Typologie"].str.startswith("T1"), "Âge BL à l'arrêté (j)"].mean(), rap.loc[rap["Typologie"].str.startswith("T1"), "Âge BL à l'arrêté (j)"].quantile(0.9), int(rap["Typologie"].str.startswith("T1").sum())],
], columns=["Indicateur", "Médiane", "Moyenne", "P90", "N"]).round(0)
delai_an = d_dn.groupby(d_dn["BL"].dt.year).agg(N=("Numéro Facture", "size"), Médiane_BL_DN=("Délai BL→DN (j)", "median"), P90_BL_DN=("Délai BL→DN (j)", lambda s: s.quantile(0.9)),
                                                 Médiane_DN_regl=("Délai DN→règlement (j)", "median")).round(0).reset_index().rename(columns={"BL": "Année BL"})
# encours vieilli (docs rapprochés avec solde)
sap["Client Name"] = sap["Client Name"].astype(str).str.strip().str.upper().str.replace(r"\s+", " ", regex=True)
ouv = sap[(sap["Clé"].notna()) & (sap["Solde abs"] > 0.01)]
aging = ouv.groupby("Tranche retard").agg(Docs=("Numéro Facture", "size"), Solde=("Solde abs", "sum")).reindex(["Non échu", "1-30 j", "31-60 j", "61-90 j", "91-180 j", "> 180 j"]).fillna(0).reset_index()
aging_cli = ouv.groupby(["Client Name", "Tranche retard"])["Solde abs"].sum().unstack(fill_value=0).reindex(columns=["Non échu", "1-30 j", "31-60 j", "61-90 j", "91-180 j", "> 180 j"], fill_value=0)
aging_cli["Total"] = aging_cli.sum(axis=1); aging_cli = aging_cli.sort_values("Total", ascending=False).reset_index()
# non démarrés par ancienneté
nd = rap[rap["Typologie"].str.startswith("T1")].copy()
nd["Tranche âge"] = pd.cut(nd["Âge BL à l'arrêté (j)"], [-1, 90, 180, 365, 10000], labels=["0-90 j", "91-180 j", "181-365 j", "> 365 j"])
nd_age = nd.groupby("Tranche âge", observed=False).agg(Clés=(KEY, "size"), Montant=(M, "sum")).reset_index()

# ------------------------------------------------------------------ 5. par client / région
def by(col):
    g = chain.groupby(col).agg(Clés=("Clé", "size"), Réclamé=("② Réclamé tracker (USD)", "sum"), Refacturé=("③ Refacturé SAP (USD)", "sum"), Encaissé=("④ Encaissé / imputé (USD)", "sum"),
                               Solde=("④ Solde client (USD)", "sum"), Exposé=("Montant exposé (USD)", "sum"),
                               T2=("Typologie de fuite", lambda s: s.str.startswith("T2").sum()), T4_T5=("Typologie de fuite", lambda s: s.str.startswith(("T4", "T5")).sum()))
    g["% refacturé"] = (g["Refacturé"] / g["Réclamé"]).replace([np.inf, np.nan], 0)
    return g.sort_values("Exposé", ascending=False).reset_index()
chain["Client"] = chain["Client"].astype(str).str.strip().str.upper().str.replace(r"\s+", " ", regex=True)
cli = by("Client"); reg = by("Région")

# ------------------------------------------------------------------ 6. référentiel navires
exp["VK"] = exp["Navire"].map(norm)
imo = exp.groupby("VK").agg(IMO=("Navir IMO", lambda s: " / ".join(sorted({str(int(x)) for x in s.dropna()}))), Noms_SAP=("Navire", lambda s: " | ".join(sorted(set(s.dropna().astype(str))))), Docs=("Numéro Facture", "size")).reset_index()
det["VK"] = det["VESSEL"].map(norm)
trk_names = det.groupby("VK").agg(Noms_tracker=("VESSEL", lambda s: " | ".join(sorted({f"« {x} »" for x in s.astype(str)}))), Lignes=("VESSEL", "size"), Variantes=("VESSEL", lambda s: len(set(s.astype(str))))).reset_index()
refnav = trk_names.merge(imo, on="VK", how="left")
refnav["IMO"] = refnav["IMO"].fillna("")
refnav["Présent SAP"] = np.where(refnav["Noms_SAP"].notna(), "Oui", "Non")
refnav["Alerte"] = np.where(refnav["Variantes"] > 1, "Plusieurs graphies dans le tracker", np.where(refnav["Présent SAP"] == "Non", "Navire absent de l'export SAP", ""))
refnav = refnav[["VK", "Noms_tracker", "Variantes", "Lignes", "Présent SAP", "Noms_SAP", "IMO", "Docs", "Alerte"]].rename(columns={"VK": "Navire normalisé", "Noms_tracker": "Graphies tracker", "Variantes": "Nb graphies", "Lignes": "Lignes tracker", "Noms_SAP": "Nom SAP", "IMO": "IMO (SAP)", "Docs": "Docs SAP (tous BL)"}).sort_values(["Alerte", "Navire normalisé"], ascending=[False, True])
# IMO partagé par plusieurs noms (SAP) → même navire sous 2 noms
imo_multi = exp.dropna(subset=["Navir IMO"]).groupby("Navir IMO")["Navire"].agg(lambda s: sorted(set(s.astype(str)))).reset_index()
imo_multi = imo_multi[imo_multi["Navire"].map(len) > 1]
imo_multi["Navire"] = imo_multi["Navire"].map(" | ".join); imo_multi["Navir IMO"] = imo_multi["Navir IMO"].astype(int).astype(str)
imo_multi.columns = ["IMO", "Noms différents dans SAP pour le même IMO"]

# ------------------------------------------------------------------ 7. doublons / réémissions SAP
sap["VK"] = sap["Navire"].map(norm)
dups = []
for (vk, bl, t), g in sap.groupby(["VK", "BL", "Ord.Type"]):
    if g["Création"].nunique() < 2: continue
    batches = g.groupby("Création").agg(n=("Numéro Facture", "size"), docs=("Numéro Facture", lambda s: " ; ".join(s)), amt=("Montant Facturée", "sum"), solde=("Solde abs", "sum"), amts=("Montant Facturée", lambda s: tuple(sorted(round(abs(x), 2) for x in s)))).reset_index()
    # réémission probable : un lot postérieur répète ≥1 montant d'un lot antérieur
    for i in range(1, len(batches)):
        prev = set().union(*[set(a) for a in batches["amts"].iloc[:i]])
        rep = [a for a in batches["amts"].iloc[i] if a in prev]
        if rep:
            dups.append([g["Navire"].iloc[0], bl.date(), t, g["Clé"].dropna().iloc[0] if g["Clé"].notna().any() else "(non suivi)", batches["Création"].iloc[i - 1].date(), batches["docs"].iloc[i - 1], round(batches["amt"].iloc[i - 1], 2), round(batches["solde"].iloc[i - 1], 2),
                         batches["Création"].iloc[i].date(), batches["docs"].iloc[i], round(batches["amt"].iloc[i], 2), len(rep), "Lot antérieur non soldé → annulation non comptabilisée ?" if batches["solde"].iloc[i - 1] > 0.01 else "Lot antérieur soldé"])
dups = pd.DataFrame(dups, columns=["Navire", "Date BL", "Type", "Clé tracker", "Lot 1 – création", "Lot 1 – docs", "Lot 1 – montant", "Lot 1 – solde", "Lot 2 – création", "Lot 2 – docs", "Lot 2 – montant", "Montants répétés", "Lecture"])

# ------------------------------------------------------------------ 8. contrôles automatiques
ctrl = []
def test(name, ok, detail): ctrl.append([name, "OK" if ok else "KO", detail])
tot_trk = round(rap[M].sum(), 2); tot_chain = round(chain["② Réclamé tracker (USD)"].sum(), 2)
test("Total tracker = total chaîne par voyage", abs(tot_trk - tot_chain) < 0.01, f"{tot_trk:,.2f} vs {tot_chain:,.2f}")
test("Somme des typologies = nombre de clés", typ["Clés"].sum() == len(rap), f"{typ['Clés'].sum()} / {len(rap)}")
docs_all = [n for v in rap["Docs SAP (n° facture)"].dropna().astype(str) for n in v.split(" ; ") if n and n != "nan"]
test("Aucun document SAP rattaché à deux clés", len(docs_all) == len(set(docs_all)), f"{len(docs_all)} rattachements, {len(set(docs_all))} documents distincts")
test("Documents rattachés tous présents dans la base SAP 2024+", set(docs_all) <= set(sap["Numéro Facture"]), f"{len(set(docs_all) - set(sap['Numéro Facture']))} manquants")
d24 = det[det["Périmètre"] == "2024+"]
test("Toutes les lignes tracker 2024+ ont une clé rapprochée", d24["Clé"].isin(rap[KEY]).all(), f"{(~d24['Clé'].isin(rap[KEY])).sum()} lignes orphelines")
test("Montant SAP abs = |ZDS| + |ZDP| sur chaque clé", (abs(rap["SAP ZDS – Débit (demurrage)"].abs() + rap["SAP ZDP – Crédit (despatch)"].abs() - rap[S]) < 0.02).all(), "")
test("Encaissé + solde = refacturé sur chaque clé", (abs(rap[E] + rap[SO] - rap[S]) < 0.02).all(), "")
paid_nodn = rap[(rap["Pay"] == "PAID") & (rap["Nb docs SAP"] == 0) & (~rap["DN/CN n° (tracker)"].astype(str).str.contains(r"9\d{7}"))]
test("Clés PAID sans doc SAP citent au moins un n° de DN/CN", len(paid_nodn) == 0, f"{len(paid_nodn)} clés PAID sans aucun n° cité : " + " ; ".join(paid_nodn["Navire"].head(10)))
neg = rap[rap[M] < 0]
test("Aucun montant tracker négatif", len(neg) == 0, f"{len(neg)}")
zero = rap[(rap[M] == 0)]
test("Montants tracker nuls (à justifier)", len(zero) == 0, f"{len(zero)} clés à 0 : " + " ; ".join(zero["Navire"].astype(str).head(10)))
dup_trk = d24.groupby("Clé").size(); multi = dup_trk[dup_trk > 1]
test("Clés tracker portées par plusieurs lignes (regroupées)", True, f"{len(multi)} clés multi-lignes (info)")
fut = sap[sap["Création"] < sap["BL"]]
test("Aucune DN/CN créée avant la date BL", len(fut) == 0, f"{len(fut)} documents : " + " ; ".join(fut["Numéro Facture"].head(10)))
test("Réémissions SAP détectées (lot antérieur non soldé)", (dups["Lecture"].str.startswith("Lot antérieur non soldé")).sum() == 0 if len(dups) else True, f"{(dups['Lecture'].str.startswith('Lot antérieur non soldé')).sum() if len(dups) else 0} cas")
test("Navires du tracker avec plusieurs graphies", (refnav["Nb graphies"] > 1).sum() == 0, f"{(refnav['Nb graphies'] > 1).sum()} navires")
test("Même IMO sous plusieurs noms dans SAP", len(imo_multi) == 0, f"{len(imo_multi)} IMO")
test("Factures de fret armateur fournies", fret is not None, "Absentes : maillon ① non mesuré" if fret is None else f"{len(fret)} lignes")
ctrl = pd.DataFrame(ctrl, columns=["Contrôle", "Résultat", "Détail"])

# ------------------------------------------------------------------ 9. échantillon d'audit
def pick(df, n, why, pieces):
    out = df.head(n).copy(); out["Motif"] = why; out["Pièces à réunir"] = pieces; return out
samp = pd.concat([
    pick(chain[chain["Typologie de fuite"].str.startswith(("T4", "T5"))].sort_values("④ Solde client (USD)", ascending=False), 5, "Solde client ouvert le plus élevé", "DN/CN SAP, relances, accusé client, avis de règlement"),
    pick(chain[chain["Typologie de fuite"].str.startswith("T2") & (chain["② Statut tracker"] == "PAID")].sort_values("② Réclamé tracker (USD)", ascending=False), 5, "Marqué PAID sans document dans l'export", "N° DN/CN cité, pièce comptable SAP, swift / avis de paiement"),
    pick(chain[chain["Typologie de fuite"].str.startswith("T3")].sort_values("Écart ②–③ (USD)", key=abs, ascending=False), 3, "Écart de montant tracker / SAP", "Calcul de demurrage (laytime, SOF), DN émise, correspondance client"),
    pick(chain[chain["Alerte doublons SAP"] == "Oui"].sort_values("③ Refacturé SAP (USD)", ascending=False), 2, "Documents SAP en surplus (réémission ?)", "Documents annulés, pièces d'annulation, lettrage"),
    pick(chain[chain["Typologie de fuite"].str.startswith("T1b")].sort_values("② Réclamé tracker (USD)", ascending=False), 3, f"Non démarré depuis > {args.delai_alerte} j", "Contrat (clause délai de réclamation), time sheet, facture de fret"),
]).reset_index(drop=True)
samp.insert(0, "N°", range(1, len(samp) + 1))
samp = samp[["N°", "Motif", "Navire", "Date BL", "Type", "Client", "② Réclamé tracker (USD)", "③ Refacturé SAP (USD)", "④ Solde client (USD)", "Typologie de fuite", "Pièces à réunir", "Clé"]]

# ------------------------------------------------------------------ 10. plan d'action chiffré
t2_amt = T.get("T2 – Réclamé, DN/CN absente de l'export SAP", 0); t45 = T.get("T4 – Refacturé, non encaissé", 0) + T.get("T5 – Refacturé, partiellement encaissé", 0)
t1b = T.get("T1b – Non démarré > délai d'alerte", 0); t1 = T.get("T1 – Non réclamé / non démarré", 0)
plan = pd.DataFrame([
    [1, "Obtenir l'export SAP complet (toutes entités / familles produit, documents annulés inclus)", "Finance / IT SAP", "Immédiat", t2_amt, f"{int(typ.loc[typ['Typologie de fuite'].str.startswith('T2'), 'Clés'].sum())} clés T2 à reclasser"],
    [2, "Obtenir les factures de fret armateur (navire, BL, demurrage payé, date)", "Logistique / Affrètement", "Immédiat", np.nan, "Maillon ① : rend l'indicateur « payé → récupéré » calculable"],
    [3, "Relancer les clients sur les débit notes ouvertes ; traiter les lots réémis non annulés", "Post-sales + Comptabilité clients", "2 semaines", t45, f"{int(typ.loc[typ['Typologie de fuite'].str.startswith(('T4', 'T5')), 'Clés'].sum())} clés ; priorité " + ", ".join(chain[chain['Typologie de fuite'].str.startswith(('T4', 'T5'))].sort_values('④ Solde client (USD)', ascending=False)['Navire'].head(2).astype(str).str.strip())],
    [4, f"Émettre les DN/CN des dossiers non démarrés > {args.delai_alerte} j (risque de prescription)", "Post-sales", "2 semaines", t1b, f"{int(typ.loc[typ['Typologie de fuite'].str.startswith('T1b'), 'Clés'].sum())} clés"],
    [5, "Corriger les écarts et les erreurs de saisie (" + ", ".join(chain[chain['Typologie de fuite'].str.startswith('T3')]['Navire'].head(3).astype(str).str.strip()) + "…)", "Post-sales", "2 semaines", float(chain.loc[chain['Typologie de fuite'].str.startswith('T3'), 'Écart ②–③ (USD)'].abs().sum()), f"{int(typ.loc[typ['Typologie de fuite'].str.startswith('T3'), 'Clés'].sum())} clés T3"],
    [6, "Mettre en place le délai cible BL → DN et le suivi mensuel (dossiers T1)", "Post-sales / Contrôle de gestion", "1 mois", t1, "Indicateur de délai, revue mensuelle"],
    [7, "Normaliser la saisie du tracker (référentiel navire/IMO, 1 BL par ligne, statuts en liste, 1 n° DN par cellule)", "Post-sales", "1 mois", np.nan, "Rend l'outil rejouable sans retraitement"],
    [8, "Régulariser les dettes envers les acheteurs (crédit notes despatch CFR, demurrage FOB) : émettre, imputer ou provisionner", "Post-sales + Comptabilité", "1 mois", sum(TD.get(k, 0) for k in ["T1 – Non réclamé / non démarré", "T1b – Non démarré > délai d'alerte", "T4 – Refacturé, non encaissé", "T5 – Refacturé, partiellement encaissé"]), "Ce n'est pas une recette : à sortir des montants récupérables"],
    [9, "Rejouer l'outil chaque mois avec les nouveaux exports", "Contrôle de gestion", "Mensuel", np.nan, "Dossier 04 - Outil rejouable"],
], columns=["#", "Action", "Responsable proposé", "Échéance", "Montant concerné (USD)", "Détail"])

# ------------------------------------------------------------------ 11. écriture du classeur
wb = Workbook(); FONT = "Arial"
hdr_fill = PatternFill("solid", fgColor="1F3864"); hdr_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
base = Font(name=FONT, size=10); bold = Font(name=FONT, size=10, bold=True); title = Font(name=FONT, size=14, bold=True, color="1F3864")
thin = Side(style="thin", color="BFBFBF"); border = Border(top=thin, bottom=thin, left=thin, right=thin)
MONEY = '#,##0.00;[Red](#,##0.00);"-"'; MONEY0 = '#,##0;[Red](#,##0);"-"'; PCT = "0.0%"; DATE = "DD/MM/YYYY"


def write_df(ws, df, r0=1, money=(), money0=(), date=(), pct=(), intc=(), widths=None, autofilter=True):
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(row=r0, column=j, value=str(c)); cell.font = hdr_font; cell.fill = hdr_fill; cell.border = border; cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, rec in enumerate(df.itertuples(index=False), r0 + 1):
        for j, v in enumerate(rec, 1):
            if isinstance(v, float) and np.isnan(v): v = None
            if isinstance(v, pd.Timestamp): v = None if pd.isna(v) else v.to_pydatetime()
            if v is pd.NaT: v = None
            if hasattr(v, "item") and not isinstance(v, (str, bytes)):
                try: v = v.item()
                except Exception: pass
            cell = ws.cell(row=i, column=j, value=v); cell.font = base; cell.border = border
            col = df.columns[j - 1]
            if col in money: cell.number_format = MONEY
            elif col in money0: cell.number_format = MONEY0
            elif col in date: cell.number_format = DATE
            elif col in pct: cell.number_format = PCT
            elif col in intc: cell.number_format = "0"
    ws.row_dimensions[r0].height = 40
    if autofilter: ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(df.columns))}{r0 + len(df)}"
    for j, c in enumerate(df.columns, 1):
        w = (widths or {}).get(c)
        if w is None:
            mx = max([len(str(c)) // 2 + 4] + [len(str(x)) for x in df[c].head(300).tolist() if x is not None and not (isinstance(x, float) and np.isnan(x))])
            w = min(max(9, mx + 2), 45)
        ws.column_dimensions[get_column_letter(j)].width = w
    return r0 + len(df) + 2


# --- Tableau de bord (formules)
ws = wb.active; ws.title = "Tableau de bord"
ws["A1"] = "Analyse experte – Demurrage / Despatch 2024-2026 : où se perd l'argent et combien est récupérable"; ws["A1"].font = title
ws["A2"] = f"Date d'arrêté (extraction SAP) : {ARRETE:%d/%m/%Y}. Délai d'alerte sans DN : {args.delai_alerte} j. Généré le {datetime.now():%d/%m/%Y}. Toutes les cellules chiffrées sont des formules sur l'onglet « Chaîne par voyage »."
ws["A2"].font = Font(name=FONT, size=9, italic=True)
CH = "'Chaîne par voyage'"; cc = {c: get_column_letter(i + 1) for i, c in enumerate(chain.columns)}; n = len(chain) + 1
rng = lambda c: f"{CH}!${cc[c]}$2:${cc[c]}${n}"
r = 4
ws.cell(row=r, column=1, value="1. Chaîne de valeur par maillon (montants USD)").font = bold; r += 1
for j, h in enumerate(["Maillon", "Montant (USD)", "Clés", "Lecture"], 1):
    c = ws.cell(row=r, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill; c.border = border
r += 1
mail = [("① Payé à l'armateur", f'=SUM({rng("① Payé armateur (USD)")})', f'=COUNT({rng("① Payé armateur (USD)")})', "Non mesuré : factures de fret absentes des fichiers reçus" if fret is None else "Factures de fret fournies"),
        ("② Réclamé (tracker post-sales)", f'=SUM({rng("② Réclamé tracker (USD)")})', f'=COUNTA({rng("Clé")})', "Montant final suivi par le post-sales"),
        ("③ Refacturé (DN/CN SAP)", f'=SUM({rng("③ Refacturé SAP (USD)")})', f'=COUNTIF({rng("③ Nb docs SAP")},">0")', "Documents ZDS/ZDP trouvés dans l'export"),
        ("④ Encaissé / imputé", f'=SUM({rng("④ Encaissé / imputé (USD)")})', "", "Montant réglé ou imputé"),
        ("    Solde client ouvert", f'=SUM({rng("④ Solde client (USD)")})', f'=COUNTIF({rng("④ Solde client (USD)")},">1")', "Refacturé non encaissé"),
        ("    Taux ② → ③ (refacturation)", f"=IF(B{r+1}=0,0,B{r+2}/B{r+1})", "", "Part du réclamé qui a un document SAP"),
        ("    Taux ③ → ④ (encaissement)", f"=IF(B{r+2}=0,0,B{r+3}/B{r+2})", "", "Part du refacturé encaissée")]
for lab, f1, f2, lec in mail:
    ws.cell(row=r, column=1, value=lab).font = bold if lab.startswith("①") or lab.startswith("②") or lab.startswith("③") or lab.startswith("④") else base
    c = ws.cell(row=r, column=2, value=f1); c.font = base; c.border = border; c.number_format = PCT if "Taux" in lab else MONEY0
    c = ws.cell(row=r, column=3, value=f2 if f2 else None); c.font = base; c.border = border; c.number_format = "0"
    ws.cell(row=r, column=4, value=lec).font = Font(name=FONT, size=9, italic=True)
    r += 1
r += 1
ws.cell(row=r, column=1, value="2. Typologie de fuite (montant exposé = réclamé si aucun doc SAP, sinon solde client)").font = bold; r += 1
for j, h in enumerate(["Typologie", "Clés", "Réclamé (USD)", "Refacturé (USD)", "Encaissé (USD)", "Montant exposé (USD)", "% du réclamé"], 1):
    c = ws.cell(row=r, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill; c.border = border
r += 1; r_typ0 = r
for tname in ["T0 – Soldé", "T1 – Non réclamé / non démarré", "T1b – Non démarré > délai d'alerte", "T2 – Réclamé, DN/CN absente de l'export SAP", "T3 – Écart de montant tracker / SAP", "T4 – Refacturé, non encaissé", "T5 – Refacturé, partiellement encaissé"]:
    ws.cell(row=r, column=1, value=tname).font = base
    crit = f'{rng("Typologie de fuite")},"{tname}"'
    for j, f in enumerate([f'=COUNTIF({crit})', f'=SUMIFS({rng("② Réclamé tracker (USD)")},{crit})', f'=SUMIFS({rng("③ Refacturé SAP (USD)")},{crit})', f'=SUMIFS({rng("④ Encaissé / imputé (USD)")},{crit})', f'=SUMIFS({rng("Montant exposé (USD)")},{crit})', f'=IF($C${r_typ0 + 7}=0,0,C{r}/$C${r_typ0 + 7})'], 2):
        c = ws.cell(row=r, column=j, value=f); c.font = base; c.border = border; c.number_format = "0" if j == 2 else (PCT if j == 7 else MONEY0)
    r += 1
ws.cell(row=r, column=1, value="Total").font = bold
for j, L in zip(range(2, 7), "BCDEF"):
    c = ws.cell(row=r, column=j, value=f"=SUM({L}{r_typ0}:{L}{r - 1})"); c.font = bold; c.border = border; c.number_format = "0" if j == 2 else MONEY0
r += 2
ws.cell(row=r, column=1, value="3. Scénarios de récupération (créances) et dettes (sens du flux : CFR demurrage = créance, CFR despatch = dette ; FOB inversé)").font = bold; r += 1
for j, h in enumerate(["Scénario", "Montant (USD)", "Définition", "Action"], 1):
    c = ws.cell(row=r, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill; c.border = border
r += 1
CR = rng("Créance exposée (USD)"); DT = rng("Dette exposée (USD)"); TY = rng("Typologie de fuite")
scen_f = {"Récupérable certain": f'=SUMIFS({CR},{TY},"T4*")+SUMIFS({CR},{TY},"T5*")',
          "Récupérable probable": f'=SUMIFS({rng("Écart ②–③ (USD)")},{TY},"T3*",{rng("Écart ②–③ (USD)")},">0",{rng("Sens du flux")},"Créance*")',
          "À confirmer (export filtré ?)": f'=SUMIFS({CR},{TY},"T2*")',
          "En cours de traitement": f'=SUMIFS({CR},{TY},"T1 *")',
          "Risque de prescription": f'=SUMIFS({CR},{TY},"T1b*")',
          "Dettes sans document dans l'export": f'=SUMIFS({DT},{TY},"T2*")',
          "Dettes à régler ou provisionner": f'=SUMIFS({DT},{TY},"T1 *")+SUMIFS({DT},{TY},"T1b*")+SUMIFS({DT},{TY},"T4*")+SUMIFS({DT},{TY},"T5*")'}
for _, s in scen.iterrows():
    ws.cell(row=r, column=1, value=s["Scénario"]).font = bold
    c = ws.cell(row=r, column=2, value=scen_f[s["Scénario"]]); c.font = base; c.border = border; c.number_format = MONEY0
    ws.cell(row=r, column=3, value=s["Définition"]).font = base; ws.cell(row=r, column=4, value=s["Action"]).font = Font(name=FONT, size=9, italic=True)
    r += 1
r += 1
ws.cell(row=r, column=1, value="4. Sens du flux par Incoterm").font = bold; r += 1
for j, h in enumerate(["Incoterm | sens", "Clés", "Réclamé (USD)", "Refacturé (USD)", "Encaissé (USD)", "Exposé (USD)"], 1):
    c = ws.cell(row=r, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill; c.border = border
r += 1
for _, srow in sens_tab.iterrows():
    lab = f"{srow['Incoterm']} | {srow['Sens du flux']}"; crit = f'{rng("Incoterm")},"{srow["Incoterm"]}",{rng("Sens du flux")},"{srow["Sens du flux"]}"'
    ws.cell(row=r, column=1, value=lab).font = base
    for j, f in enumerate([f'=COUNTIFS({crit})', f'=SUMIFS({rng("② Réclamé tracker (USD)")},{crit})', f'=SUMIFS({rng("③ Refacturé SAP (USD)")},{crit})', f'=SUMIFS({rng("④ Encaissé / imputé (USD)")},{crit})', f'=SUMIFS({rng("Montant exposé (USD)")},{crit})'], 2):
        c = ws.cell(row=r, column=j, value=f); c.font = base; c.border = border; c.number_format = "0" if j == 2 else MONEY0
    r += 1
r += 1
ws.cell(row=r, column=1, value="5. Alertes").font = bold; r += 1
for lab, f in [("Clés avec documents SAP en surplus (réémission ?)", f'=COUNTIF({rng("Alerte doublons SAP")},"Oui")'),
               (f"Clés sans document SAP depuis plus de {args.delai_alerte} jours (T1b + T2 anciennes)", f'=COUNTIF({rng("Alerte délai")},"*sans DN")'),
               ("Clés avec retard > 90 j sur échéance", f'=COUNTIF({rng("④ Tranche retard")},"91-180 j")+COUNTIF({rng("④ Tranche retard")},"> 180 j")'),
               ("Contrôles automatiques en échec (onglet Contrôles)", f"=COUNTIF(Contrôles!$B$2:$B${len(ctrl) + 1},\"KO\")")]:
    ws.cell(row=r, column=1, value=lab).font = base
    c = ws.cell(row=r, column=2, value=f); c.font = base; c.border = border; c.number_format = "0"; r += 1
ws.column_dimensions["A"].width = 52; ws.column_dimensions["B"].width = 18; ws.column_dimensions["C"].width = 16
for L in "DEFG": ws.column_dimensions[L].width = 18
ws.column_dimensions["D"].width = 60
ws.sheet_view.showGridLines = False

# --- Chaîne par voyage
ws2 = wb.create_sheet("Chaîne par voyage")
write_df(ws2, chain, money=[c for c in chain.columns if "(USD)" in c], date=["Date BL", "① Date paiement fret", "③ 1ère DN émise", "④ 1er règlement", "④ Échéance"],
         intc=["Année", "③ Nb docs SAP", "③ Délai BL→DN (j)", "④ Délai DN→règlement (j)", "④ Retard vs échéance (j)", "Âge BL à l'arrêté (j)"], widths={"Clé": 36, "Typologie de fuite": 40, "Client": 30, "② DN/CN citées": 30, "Méthode": 30})
ws2.freeze_panes = "C2"
tc = cc["Typologie de fuite"]; last = get_column_letter(len(chain.columns))
for pat, col in [("T2", "F8CBAD"), ("T4", "FFE699"), ("T5", "FFF2CC"), ("T3", "FCE4D6"), ("T1b", "E7E6E6"), ("T0", "E2EFDA")]:
    ws2.conditional_formatting.add(f"A2:{last}{n}", FormulaRule(formula=[f'LEFT(${tc}2,{len(pat)})="{pat}"'], fill=PatternFill("solid", fgColor=col)))

# --- Typologie & scénarios (valeurs calculées, contrôle des formules)
ws3 = wb.create_sheet("Typologie (valeurs)")
r = write_df(ws3, typ.rename(columns={"Réclamé": "Réclamé (USD)", "Refacturé": "Refacturé (USD)", "Encaissé": "Encaissé (USD)", "Exposé": "Montant exposé (USD)", "Créance": "dont créances (USD)", "Dette": "dont dettes (USD)"}), money0=["Réclamé (USD)", "Refacturé (USD)", "Encaissé (USD)", "Montant exposé (USD)", "dont créances (USD)", "dont dettes (USD)"], widths={"Typologie de fuite": 44})
ws3.cell(row=r, column=1, value="Sens du flux par Incoterm (valeurs)").font = bold
r = write_df(ws3, sens_tab.rename(columns={"Réclamé": "Réclamé (USD)", "Refacturé": "Refacturé (USD)", "Encaissé": "Encaissé (USD)", "Exposé": "Exposé (USD)"}), r0=r + 1, money0=["Réclamé (USD)", "Refacturé (USD)", "Encaissé (USD)", "Exposé (USD)"], widths={"Sens du flux": 36}, autofilter=False)
ws3.cell(row=r, column=1, value="Scénarios de récupération (valeurs)").font = bold
write_df(ws3, scen, r0=r + 1, money0=["Montant (USD)"], widths={"Définition": 50, "Action": 60}, autofilter=False)
ws3.cell(row=1, column=8, value="Onglet de contrôle : mêmes agrégats calculés hors formules ; doit être égal au Tableau de bord.").font = Font(name=FONT, size=9, italic=True)

# --- Délais & encours
ws4 = wb.create_sheet("Délais & encours vieilli")
ws4["A1"] = "Délais de la chaîne (jours) — documents SAP rapprochés"; ws4["A1"].font = title
r = write_df(ws4, delais, r0=3, intc=["Médiane", "Moyenne", "P90", "N"], widths={"Indicateur": 48}, autofilter=False)
ws4.cell(row=r, column=1, value="Délais par année de BL").font = bold
r = write_df(ws4, delai_an, r0=r + 1, intc=list(delai_an.columns[1:]), autofilter=False)
ws4.cell(row=r, column=1, value=f"Encours vieilli des DN/CN rapprochées avec solde (retard calculé à la date d'arrêté {ARRETE:%d/%m/%Y} par rapport à l'échéance SAP)").font = bold
r = write_df(ws4, aging.rename(columns={"Solde": "Solde (USD)"}), r0=r + 1, money0=["Solde (USD)"], intc=["Docs"], autofilter=False)
ws4.cell(row=r, column=1, value="Encours vieilli par client (USD)").font = bold
r = write_df(ws4, aging_cli, r0=r + 1, money0=[c for c in aging_cli.columns if c != "Client Name"], widths={"Client Name": 40}, autofilter=False)
ws4.cell(row=r, column=1, value="Dossiers non démarrés (T1) par ancienneté depuis la date BL").font = bold
write_df(ws4, nd_age.rename(columns={"Montant": "Montant (USD)"}), r0=r + 1, money0=["Montant (USD)"], intc=["Clés"], autofilter=False)

# --- Par client / région
ws5 = wb.create_sheet("Par client")
cli2 = cli.rename(columns={"Réclamé": "Réclamé (USD)", "Refacturé": "Refacturé (USD)", "Encaissé": "Encaissé (USD)", "Solde": "Solde (USD)", "Exposé": "Exposé (USD)", "T2": "Clés T2 (DN absente)", "T4_T5": "Clés T4/T5 (non encaissé)"})
write_df(ws5, cli2, money0=[c for c in cli2.columns if "(USD)" in c], pct=["% refacturé"], intc=["Clés", "Clés T2 (DN absente)", "Clés T4/T5 (non encaissé)"], widths={"Client": 42})
ws6 = wb.create_sheet("Par région")
reg2 = reg.rename(columns=dict(zip(cli.columns, cli2.columns)))
write_df(ws6, reg2, money0=[c for c in reg2.columns if "(USD)" in c], pct=["% refacturé"], intc=["Clés", "Clés T2 (DN absente)", "Clés T4/T5 (non encaissé)"], widths={"Région": 40})

# --- Référentiel navires
ws7 = wb.create_sheet("Référentiel navires")
r = write_df(ws7, refnav, intc=["Nb graphies", "Lignes tracker", "Docs SAP (tous BL)"], widths={"Graphies tracker": 40, "Nom SAP": 30, "Alerte": 36})
ws7.cell(row=r, column=1, value="IMO portant plusieurs noms dans SAP (même navire, graphies différentes)").font = bold
write_df(ws7, imo_multi, r0=r + 1, widths={"Noms différents dans SAP pour le même IMO": 60}, autofilter=False)

# --- Doublons SAP
ws8 = wb.create_sheet("Réémissions SAP")
ws8["A1"] = "Lots de documents SAP créés à des dates différentes sur un même navire / BL / type, dont un lot postérieur répète des montants d'un lot antérieur (annulation puis réémission probable). Base SAP 2024+, suivie ou non par le tracker."
ws8["A1"].font = Font(name=FONT, size=9, italic=True)
write_df(ws8, dups, r0=3, money=["Lot 1 – montant", "Lot 1 – solde", "Lot 2 – montant"], date=["Date BL", "Lot 1 – création", "Lot 2 – création"], intc=["Montants répétés"], widths={"Lot 1 – docs": 40, "Lot 2 – docs": 40, "Lecture": 48, "Clé tracker": 34})

# --- Contrôles
ws9 = wb.create_sheet("Contrôles")
write_df(ws9, ctrl, widths={"Contrôle": 60, "Détail": 80}, autofilter=False)
ws9.conditional_formatting.add(f"B2:B{len(ctrl) + 1}", FormulaRule(formula=['$B2="KO"'], fill=PatternFill("solid", fgColor="F8CBAD")))
ws9.conditional_formatting.add(f"B2:B{len(ctrl) + 1}", FormulaRule(formula=['$B2="OK"'], fill=PatternFill("solid", fgColor="E2EFDA")))

# --- Échantillon d'audit
ws10 = wb.create_sheet("Échantillon audit")
ws10["A1"] = f"{len(samp)} dossiers à auditer sur pièces pour calibrer le taux d'erreur du rapprochement automatique. Colonnes « Conclusion audit » et « Écart constaté » à remplir."; ws10["A1"].font = Font(name=FONT, size=9, italic=True)
samp["Conclusion audit"] = ""; samp["Écart constaté (USD)"] = None
write_df(ws10, samp, r0=3, money=["② Réclamé tracker (USD)", "③ Refacturé SAP (USD)", "④ Solde client (USD)", "Écart constaté (USD)"], date=["Date BL"], intc=["N°"], widths={"Motif": 36, "Pièces à réunir": 50, "Typologie de fuite": 40, "Client": 30, "Conclusion audit": 30, "Clé": 34})
for i in range(4, 4 + len(samp)):
    for col in ("Conclusion audit", "Écart constaté (USD)"):
        ws10.cell(row=i, column=list(samp.columns).index(col) + 1).fill = PatternFill("solid", fgColor="FFFF00")

# --- Plan d'action
ws11 = wb.create_sheet("Plan d'action chiffré")
write_df(ws11, plan, money0=["Montant concerné (USD)"], intc=["#"], widths={"Action": 70, "Détail": 55, "Responsable proposé": 30}, autofilter=False)

# --- Hypothèses
ws12 = wb.create_sheet("Hypothèses & paramètres")
hyp = [("Paramètres", ""), ("Date d'arrêté", f"{ARRETE:%d/%m/%Y} (Report Run Date de l'export SAP) : tous les retards et âges sont calculés à cette date."),
       ("Délai d'alerte sans DN", f"{args.delai_alerte} jours depuis la date BL (paramètre --delai-alerte). À aligner sur la clause contractuelle de délai de réclamation."),
       ("Tolérance d'écart", f"{args.tolerance} USD ou 0,5 % (héritée du rapprochement)."), ("", ""),
       ("Typologies", ""), ("T0", "Soldé : documents SAP trouvés, solde client nul."), ("T1 / T1b", "Non démarré (statut tracker NOT YET STARTED, aucun doc SAP) ; T1b si l'âge depuis BL dépasse le délai d'alerte."),
       ("T2", "Réclamé (PAID / PENDING) mais aucun document ZDS/ZDP dans l'export : à confirmer avec un export complet avant toute conclusion."),
       ("T3", "Écart de montant entre tracker et SAP au-delà de la tolérance."), ("T4 / T5", "Refacturé, non encaissé / partiellement encaissé (solde client > tolérance)."),
       ("Montant exposé", "Réclamé tracker si aucun document SAP ; sinon solde client restant."),
       ("Sens du flux", "CFR / CIF / DPU : demurrage = créance sur l'acheteur (débit note ZDS), despatch = dette envers l'acheteur (crédit note ZDP). FOB : inversé — le demurrage au port de chargement est dû par le vendeur à l'acheteur (paiement OP / crédit note), le despatch est une débit note à l'acheteur. Les scénarios de récupération ne retiennent que les créances ; les dettes sont isolées."), ("", ""),
       ("Limites", ""), ("Maillon ①", "Payé à l'armateur : non mesuré, factures de fret absentes. Le modèle 'modele_factures_fret.xlsx' décrit le format attendu ; une fois fourni (--fret), la colonne ① se remplit et l'écart ①–② devient calculable."),
       ("Encaissé", "|Montant| − |Solde client| : un document annulé non lettré apparaît non encaissé ; voir onglet Réémissions SAP."),
       ("Délais", "Date de création de la DN/CN SAP = date d'émission ; Date Reglement 1 = premier règlement. Les délais négatifs ou nuls sont conservés tels quels."),
       ("Référentiel", "Normalisation des noms de navires (majuscules, ponctuation, espaces). L'IMO SAP permet d'unifier les graphies ; le tracker ne porte pas d'IMO.")]
ws12["A1"] = "Hypothèses, paramètres et limites"; ws12["A1"].font = title
for i, (k, v) in enumerate(hyp, 3):
    a = ws12.cell(row=i, column=1, value=k); a.font = bold if v == "" else base
    b = ws12.cell(row=i, column=2, value=v); b.font = base; b.alignment = Alignment(wrap_text=True, vertical="top")
ws12.column_dimensions["A"].width = 26; ws12.column_dimensions["B"].width = 130

wb.calculation.fullCalcOnLoad = True
wb.save(args.out)
summary = dict(arrete=str(ARRETE.date()), n=len(chain), typologie=typ.to_dict("records"), scenarios=scen.to_dict("records"), sens=sens_tab.to_dict("records"), scen_sens=scen_sens, delais=delais.to_dict("records"), aging=aging.to_dict("records"),
               nd_age=nd_age.astype(str).to_dict("records"), controles=ctrl.to_dict("records"), dups=len(dups), dups_nonsolde=int(dups["Lecture"].str.startswith("Lot antérieur non soldé").sum()) if len(dups) else 0,
               refnav_multi=int((refnav["Nb graphies"] > 1).sum()), refnav_absent=int((refnav["Présent SAP"] == "Non").sum()), imo_multi=len(imo_multi), sample=len(samp), plan=plan.astype(str).to_dict("records"),
               top_clients=cli2.head(10).round(0).astype(str).to_dict("records"), delai_an=delai_an.astype(str).to_dict("records"), aging_cli=aging_cli.head(8).round(0).astype(str).to_dict("records"))
def _clean(o):
    if isinstance(o, dict): return {k: _clean(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)): return [_clean(v) for v in o]
    if isinstance(o, float) and np.isnan(o): return None
    if o is pd.NaT or (hasattr(o, "__class__") and o.__class__.__name__ in ("NAType", "NaTType")): return None
    if hasattr(o, "item") and not isinstance(o, (str, bytes)):
        try: return _clean(o.item())
        except Exception: return str(o)
    if isinstance(o, (pd.Timestamp, datetime)): return str(o)
    return o
json.dump(_clean(summary), open(args.out.replace(".xlsx", "_summary.json"), "w"), ensure_ascii=False, indent=1, default=str, allow_nan=False)
print("saved", args.out)
print(typ.round(0).to_string()); print(scen[["Scénario", "Montant (USD)"]].round(0).to_string()); print(delais.to_string()); print(aging.round(0).to_string()); print(ctrl.to_string())
print("dups", len(dups), "refnav multi", (refnav["Nb graphies"] > 1).sum(), "absent SAP", (refnav["Présent SAP"] == "Non").sum(), "imo multi", len(imo_multi), "sample", len(samp))
