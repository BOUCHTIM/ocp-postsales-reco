"""Suivi des crédits documentaires (LC) : échéances, dates limites d'expédition, présentation des documents, réserves (discrepancies), frais bancaires, délais.
Entrée : --lc registre LC : N° LC, Client, Banque émettrice, Montant, Devise, Date d'émission, Date limite d'expédition, Date d'expiration, Navire, Date BL,
         Date présentation documents, Date acceptation / paiement, Réserves (texte ou nombre), Frais bancaires, Statut [, Incoterm, N° facture]
Sortie : classeur « Suivi LC » : alertes (expiration sous N jours, BL après date limite, documents non présentés, réserves), délais BL → présentation → paiement, frais par banque / client, synthèse en formules.
Usage : lc.py --lc registre.xlsx --out lc.xlsx [--arrete AAAA-MM-JJ] [--alerte 15] [--config alias.json]"""
import argparse, json
import pandas as pd, numpy as np
from openpyxl import Workbook
from xlsx_common import *

ap = argparse.ArgumentParser(); ap.add_argument("--lc", required=True); ap.add_argument("--out", required=True); ap.add_argument("--arrete"); ap.add_argument("--alerte", type=int, default=15, help="jours avant expiration déclenchant l'alerte"); ap.add_argument("--config")
a = ap.parse_args()
DEF = {"lc": {"N° LC": ["LC NUMBER", "LC NO", "L/C", "DOCUMENTARY CREDIT", "N° CREDOC", "CREDOC"], "Client": ["CUSTOMER", "APPLICANT", "BUYER", "CLIENT"], "Banque émettrice": ["ISSUING BANK", "BANK", "BANQUE"], "Banque notificatrice": ["ADVISING BANK", "CONFIRMING BANK"],
              "Montant": ["AMOUNT", "LC AMOUNT", "MONTANT", "VALUE"], "Devise": ["CURRENCY", "DEVISE"], "Date d'émission": ["ISSUE DATE", "DATE OF ISSUE", "DATE ÉMISSION", "ISSUED"], "Date limite d'expédition": ["LATEST SHIPMENT DATE", "LATEST SHIPMENT", "DATE LIMITE EXPÉDITION", "LSD"],
              "Date d'expiration": ["EXPIRY DATE", "EXPIRY", "DATE EXPIRATION", "VALIDITY"], "Navire": ["VESSEL", "NAVIRE"], "Date BL": ["B/L DATE", "BL DATE", "DATE BL", "SHIPMENT DATE"], "Date présentation documents": ["PRESENTATION DATE", "DOCS PRESENTED", "DATE PRÉSENTATION", "NEGOTIATION DATE"],
              "Date acceptation / paiement": ["PAYMENT DATE", "ACCEPTANCE DATE", "DATE PAIEMENT", "MATURITY PAID", "SETTLEMENT DATE"], "Réserves": ["DISCREPANCIES", "DISCREPANCY", "RÉSERVES", "RESERVES"], "Frais bancaires": ["BANK CHARGES", "FEES", "FRAIS", "CHARGES"],
              "Statut": ["STATUS", "STATUT"], "Incoterm": ["INCOTERM", "INCOTERMS"], "N° facture": ["INVOICE", "BILLING DOCUMENT", "FACTURE"], "Tenor (j)": ["TENOR", "USANCE", "PAYMENT TERMS", "DÉLAI DE PAIEMENT"]}}
cfg = load_config(a.config, DEF)
L = read_any(a.lc, aliases=cfg["lc"], required=["N° LC", "Client", "Montant"])
ARRETE = pd.Timestamp(a.arrete) if a.arrete else pd.Timestamp.today().normalize()
for c in ["Date d'émission", "Date limite d'expédition", "Date d'expiration", "Date BL", "Date présentation documents", "Date acceptation / paiement"]:
    if c in L.columns: L[c] = L[c].map(to_date)
    else: L[c] = pd.NaT
for c in ["Montant", "Frais bancaires"]:
    if c in L.columns: L[c] = L[c].map(to_num)
    else: L[c] = np.nan
L["Statut"] = L["Statut"].astype(str).str.strip().str.upper() if "Statut" in L.columns else ""
L["Réserves"] = L["Réserves"] if "Réserves" in L.columns else ""
L["Nb réserves"] = L["Réserves"].map(lambda v: 0 if pd.isna(v) or str(v).strip() in ("", "0", "nan", "NONE", "NIL") else (int(float(v)) if str(v).replace(".", "", 1).isdigit() else len([x for x in str(v).split(";") if x.strip()])))
paid = L["Date acceptation / paiement"].notna() | L["Statut"].isin(["PAID", "PAYÉ", "SETTLED", "CLOSED", "CLÔTURÉ", "RÉGLÉ"])
L["Jours avant expiration"] = (L["Date d'expiration"] - ARRETE).dt.days
L["Délai BL → présentation (j)"] = (L["Date présentation documents"] - L["Date BL"]).dt.days
L["Délai présentation → paiement (j)"] = (L["Date acceptation / paiement"] - L["Date présentation documents"]).dt.days
L["Délai BL → paiement (j)"] = (L["Date acceptation / paiement"] - L["Date BL"]).dt.days
L["Expédition après date limite"] = np.where(L["Date BL"].notna() & L["Date limite d'expédition"].notna() & (L["Date BL"] > L["Date limite d'expédition"]), "Oui", "")
L["Présentation après expiration"] = np.where(L["Date présentation documents"].notna() & L["Date d'expiration"].notna() & (L["Date présentation documents"] > L["Date d'expiration"]), "Oui", "")


def alerte(r):
    if paid[r.name]: return "Réglé"
    if pd.notna(r["Date d'expiration"]) and r["Jours avant expiration"] < 0 and pd.isna(r["Date présentation documents"]): return "LC expirée, documents non présentés"
    if r["Expédition après date limite"] == "Oui": return "BL postérieur à la date limite d'expédition (réserve probable)"
    if r["Présentation après expiration"] == "Oui": return "Documents présentés après expiration"
    if pd.notna(r["Date BL"]) and pd.isna(r["Date présentation documents"]) and (ARRETE - r["Date BL"]).days > 21: return "Documents non présentés > 21 j après BL (délai UCP 600)"
    if r["Nb réserves"] > 0 and pd.isna(r["Date acceptation / paiement"]): return "Réserves en cours, paiement non reçu"
    if pd.notna(r["Date d'expiration"]) and 0 <= r["Jours avant expiration"] <= a.alerte: return f"Expire sous {a.alerte} j"
    if pd.notna(r["Date présentation documents"]) and pd.isna(r["Date acceptation / paiement"]) and (ARRETE - r["Date présentation documents"]).days > 10: return "Paiement attendu (> 10 j après présentation)"
    return "" if pd.notna(r["Date BL"]) else "LC ouverte, expédition à venir"


L["Alerte"] = L.apply(alerte, axis=1)
L["Statut calculé"] = np.where(paid, "Réglée", np.where(L["Date présentation documents"].notna(), "Documents présentés", np.where(L["Date BL"].notna(), "Expédiée, documents à présenter", "Ouverte, non expédiée")))
L["Montant non réglé"] = np.where(paid, 0.0, L["Montant"])
cols = [c for c in ["N° LC", "Client", "Banque émettrice", "Banque notificatrice", "Montant", "Devise", "Montant non réglé", "Date d'émission", "Date limite d'expédition", "Date d'expiration", "Jours avant expiration", "Navire", "Date BL", "Incoterm", "N° facture", "Date présentation documents", "Date acceptation / paiement", "Tenor (j)", "Réserves", "Nb réserves", "Frais bancaires", "Délai BL → présentation (j)", "Délai présentation → paiement (j)", "Délai BL → paiement (j)", "Expédition après date limite", "Présentation après expiration", "Statut", "Statut calculé", "Alerte"] if c in L.columns]
L = L[cols].sort_values(["Alerte", "Date d'expiration"], ascending=[False, True]).reset_index(drop=True)
wb = Workbook(); ws = wb.active; ws.title = "Synthèse"
ws["A1"] = "Suivi des crédits documentaires (LC)"; ws["A1"].font = TITLE
ws["A2"] = f"Arrêté : {ARRETE:%d/%m/%Y}. Alertes calculées : expiration, date limite d'expédition, présentation (21 j UCP 600), réserves, paiement attendu. Formules sur l'onglet « LC »."; ws["A2"].font = NOTE
cm = colmap(L); n = len(L); R_ = lambda c: rng("LC", cm[c], n)
r = kpi_rows(ws, 4, [("LC suivies", f'=COUNTA({R_("N° LC")})', "0"), ("Montant total", f'=SUM({R_("Montant")})', MONEY0), ("Montant non réglé", f'=SUM({R_("Montant non réglé")})', MONEY0), ("Frais bancaires cumulés", f'=SUM({R_("Frais bancaires")})', MONEY0),
                    ("Délai médian BL → présentation (j)", f'=MEDIAN({R_("Délai BL → présentation (j)")})', "0"), ("Délai médian présentation → paiement (j)", f'=MEDIAN({R_("Délai présentation → paiement (j)")})', "0"), ("Délai médian BL → paiement (j)", f'=MEDIAN({R_("Délai BL → paiement (j)")})', "0"),
                    ("LC avec réserves", f'=COUNTIF({R_("Nb réserves")},">0")', "0")], "1. Volumes, montants, délais")
alerts = ["LC expirée, documents non présentés", "BL postérieur à la date limite d'expédition (réserve probable)", "Documents présentés après expiration", "Documents non présentés > 21 j après BL (délai UCP 600)", "Réserves en cours, paiement non reçu", f"Expire sous {a.alerte} j", "Paiement attendu (> 10 j après présentation)", "LC ouverte, expédition à venir"]
r = kpi_rows(ws, r, [(al, f'=COUNTIF({R_("Alerte")},"{al}")', "0") for al in alerts] + [("Montant exposé (alertes hors « ouverte »)", f'=SUMIFS({R_("Montant non réglé")},{R_("Alerte")},"<>")-SUMIFS({R_("Montant non réglé")},{R_("Alerte")},"LC ouverte, expédition à venir")-SUMIFS({R_("Montant non réglé")},{R_("Alerte")},"Réglé")', MONEY0)], "2. Alertes")
r = kpi_rows(ws, r, [(s, f'=COUNTIF({R_("Statut calculé")},"{s}")', "0") for s in ["Ouverte, non expédiée", "Expédiée, documents à présenter", "Documents présentés", "Réglée"]], "3. Statut")
ws.column_dimensions["A"].width = 60; ws.column_dimensions["B"].width = 18; ws.sheet_view.showGridLines = False
ws2 = wb.create_sheet("LC"); write_df(ws2, L, money=["Montant", "Montant non réglé", "Frais bancaires"], date=[c for c in L.columns if c.startswith("Date")], intc=[c for c in L.columns if "(j)" in c], widths={"Alerte": 44, "Réserves": 36, "Client": 30})
last = get_column_letter(len(L.columns)); ac = cm["Alerte"]
ws2.conditional_formatting.add(f"A2:{last}{n + 1}", FormulaRule(formula=[f'OR(LEFT(${ac}2,10)="LC expirée",LEFT(${ac}2,12)="BL postérieur",LEFT(${ac}2,20)="Documents présentés a")'], fill=FILL_RED))
ws2.conditional_formatting.add(f"A2:{last}{n + 1}", FormulaRule(formula=[f'OR(LEFT(${ac}2,6)="Expire",LEFT(${ac}2,8)="Réserves",LEFT(${ac}2,9)="Documents",LEFT(${ac}2,8)="Paiement")'], fill=FILL_YEL))
ws2.conditional_formatting.add(f"A2:{last}{n + 1}", FormulaRule(formula=[f'${ac}2="Réglé"'], fill=FILL_GRN))
bk = L.groupby("Banque émettrice").agg(LC=("N° LC", "size"), Montant=("Montant", "sum"), Frais=("Frais bancaires", "sum"), Réserves=("Nb réserves", "sum"), Délai_paiement_médian=("Délai présentation → paiement (j)", "median")).reset_index() if "Banque émettrice" in L.columns else pd.DataFrame()
if len(bk): write_df(wb.create_sheet("Par banque"), bk, money0=["Montant", "Frais"], intc=["LC", "Réserves", "Délai_paiement_médian"], autofilter=False)
cl = L.groupby("Client").agg(LC=("N° LC", "size"), Montant=("Montant", "sum"), Non_réglé=("Montant non réglé", "sum"), Frais=("Frais bancaires", "sum"), Réserves=("Nb réserves", "sum"), Alertes=("Alerte", lambda s: ((s != "") & (s != "Réglé") & (s != "LC ouverte, expédition à venir")).sum())).sort_values("Non_réglé", ascending=False).reset_index()
write_df(wb.create_sheet("Par client"), cl, money0=["Montant", "Non_réglé", "Frais"], intc=["LC", "Réserves", "Alertes"], autofilter=False)
summary = {"module": "lc", "arrete": str(ARRETE.date()), "lc": n, "montant": float(L["Montant"].sum()), "non_regle": float(L["Montant non réglé"].sum()), "frais": float(L["Frais bancaires"].sum()), "alertes": L["Alerte"].value_counts().to_dict(), "statuts": L["Statut calculé"].value_counts().to_dict(),
           "delai_bl_presentation_median": float(L["Délai BL → présentation (j)"].median()) if L["Délai BL → présentation (j)"].notna().any() else None, "delai_presentation_paiement_median": float(L["Délai présentation → paiement (j)"].median()) if L["Délai présentation → paiement (j)"].notna().any() else None,
           "reserves": int((L["Nb réserves"] > 0).sum()), "par_client": cl.head(10).to_dict("records")}
finish(wb, a.out, summary)
print(json.dumps({k: summary[k] for k in ["lc", "montant", "non_regle", "frais", "alertes"]}, ensure_ascii=False, default=str))
