"""Génère les modèles Excel (colonnes attendues + une ligne d'exemple en bleu + légende) pour chaque module dans assets/modeles/.
Usage : make_templates.py [dossier de sortie]"""
import sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "assets", "modeles")
os.makedirs(OUT, exist_ok=True)
T = {
 "modele_o2c.xlsx": {
  "Commandes (VA05)": (["N° commande", "Client", "Contrat", "Produit", "Quantité commandée", "Incoterm", "Prix unitaire", "Devise", "Date commande"], ["60001234", "CLIENT A", "CTR-2026-001", "DAP STANDARD", 30000, "CFR", 520, "USD", "2026-03-01"]),
  "Livraisons (VL06 - BL)": (["N° livraison", "N° commande", "Navire", "Date BL", "Quantité livrée", "Port"], ["80004567", "60001234", "ALPHA STAR", "2026-03-20", 29850, "Jorf Lasfar"]),
  "Factures (VF05)": (["N° facture", "N° livraison", "N° commande", "Date facture", "Montant", "Échéance", "Client", "Devise"], ["90012345", "80004567", "60001234", "2026-03-25", 15522000, "2026-06-23", "CLIENT A", "USD"]),
  "Règlements (FBL5N)": (["N° facture", "Montant réglé", "Date règlement", "Solde"], ["90012345", 15522000, "2026-06-20", 0])},
 "modele_balance_agee.xlsx": {
  "Postes ouverts (FBL5N)": (["Client", "N° document", "Type", "Date document", "Échéance", "Montant", "Devise", "Référence", "Navire", "Texte"], ["CLIENT A", "90012345", "Facture", "2026-03-25", "2026-06-23", 15522000, "USD", "ALPHA STAR BL 20/03/2026", "ALPHA STAR", ""]),
  "CA mensuel (optionnel)": (["Mois", "Montant", "Client"], ["2026-03", 48000000, ""])},
 "modele_lc.xlsx": {
  "Registre LC": (["N° LC", "Client", "Banque émettrice", "Banque notificatrice", "Montant", "Devise", "Date d'émission", "Date limite d'expédition", "Date d'expiration", "Navire", "Date BL", "Incoterm", "N° facture", "Date présentation documents", "Date acceptation / paiement", "Tenor (j)", "Réserves", "Frais bancaires", "Statut"],
                   ["LC-2026-0042", "CLIENT A", "BANK OF X", "ATTIJARI", 15522000, "USD", "2026-02-10", "2026-03-31", "2026-04-21", "ALPHA STAR", "2026-03-20", "CFR", "90012345", "2026-04-02", "2026-06-20", 90, "", 4200, "PAID"])},
 "modele_reclamations.xlsx": {
  "Registre": (["N° réclamation", "Client", "Navire", "Date BL", "Produit", "Type", "Cause", "Date réception", "Montant réclamé", "Montant accepté", "Statut", "Date clôture", "Gestionnaire", "Action", "Commentaire"],
               ["CLM-2026-017", "CLIENT A", "ALPHA STAR", "2026-03-20", "DAP STANDARD", "Qualité", "Humidité hors spécification", "2026-04-05", 85000, 40000, "CLOSED", "2026-05-02", "S. ALAMI", "Crédit note partielle", ""])},
 "modele_laytime.xlsx": {
  "Escales (SOF)": (["Navire", "Date BL", "Port", "Opération", "Incoterm", "Quantité", "Cadence (t/jour)", "Laytime autorisé (h)", "NOR tendue", "Début comptage", "Fin opérations", "Temps utilisé (h)", "Exceptions (h)", "Taux demurrage (USD/jour)", "Taux despatch (USD/jour)", "Montant payé/réclamé", "Source"],
                    ["ALPHA STAR", "2026-03-20", "Jorf Lasfar", "Chargement", "CFR", 29850, 10000, "", "2026-03-16 08:00", "2026-03-16 14:00", "2026-03-20 06:00", "", 6, 25000, 12500, 48000, "tracker"])},
}
LEG = {"modele_o2c.xlsx": "Une ligne par commande / livraison / facture / règlement. Les n° de commande et de livraison font la jointure ; la ligne bleue est un exemple à supprimer.",
       "modele_balance_agee.xlsx": "Postes ouverts clients à la date d'arrêté (montant signé : + créance, − avoir ou acompte). Le CA mensuel sert au DSO ; sans lui, le DSO est approximé sur les factures ouvertes.",
       "modele_lc.xlsx": "Une ligne par crédit documentaire. Dates au format jj/mm/aaaa ou aaaa-mm-jj. « Réserves » : nombre ou liste séparée par des points-virgules.",
       "modele_reclamations.xlsx": "Une ligne par réclamation. Type libre (Qualité, Quantité, Documentation, Délai, Litige commercial…). Gestionnaire = personne en charge (sert à la charge d'équipe).",
       "modele_laytime.xlsx": "Une ligne par escale. Fournir soit Laytime autorisé (h), soit Quantité + Cadence. Fournir soit Début/Fin (date-heure), soit Temps utilisé (h). Exceptions = heures non comptées (pluie, fériés SHEX, pannes)."}
for fname, sheets in T.items():
    wb = Workbook(); wb.remove(wb.active)
    for sh, (cols, ex) in sheets.items():
        ws = wb.create_sheet(sh)
        for j, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=j, value=c); cell.font = Font(name="Arial", bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F3864"); ws.column_dimensions[get_column_letter(j)].width = max(12, min(28, len(c) + 4))
        for j, v in enumerate(ex, 1): ws.cell(row=2, column=j, value=v).font = Font(name="Arial", color="0000FF")
        ws.cell(row=4, column=1, value="Légende : " + LEG[fname]).font = Font(name="Arial", italic=True, size=9)
    wb.save(os.path.join(OUT, fname)); print("modèle", fname)
