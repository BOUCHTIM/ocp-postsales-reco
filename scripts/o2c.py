"""Order-to-cash par cargaison : contrat / commande → livraison (BL) → facture → encaissement.
Entrées (xlsx/csv, alias configurables) :
  --orders     commandes de vente (ex. SAP VA05) : N° commande, Client, Contrat, Produit, Quantité commandée, Incoterm, Prix unitaire, Devise
  --deliveries livraisons / connaissements (ex. VL06, SOF) : N° livraison, N° commande, Navire, Date BL, Quantité livrée, Port
  --invoices   factures (ex. VF05) : N° facture, N° livraison ou N° commande, Date facture, Montant, Échéance, Client
  --payments   postes clients (ex. FBL5N) : N° facture (référence), Montant réglé, Date règlement, Solde   [optionnel]
Sortie : un classeur « Order-to-cash » (une ligne par cargaison navire + BL) : statut, montants, quantités, délais, alertes ; synthèse en formules.
Usage : o2c.py --orders A --deliveries B --invoices C [--payments D] --out o2c.xlsx [--config alias.json] [--arrete AAAA-MM-JJ]"""
import argparse, json
import pandas as pd, numpy as np
from openpyxl import Workbook
from xlsx_common import *

ap = argparse.ArgumentParser(); ap.add_argument("--orders", required=True); ap.add_argument("--deliveries", required=True); ap.add_argument("--invoices", required=True); ap.add_argument("--payments")
ap.add_argument("--out", required=True); ap.add_argument("--config"); ap.add_argument("--arrete"); ap.add_argument("--delai-facture", type=int, default=15, help="jours BL → facture au-delà desquels on alerte")
a = ap.parse_args()
DEF = {"orders": {"N° commande": ["SALES DOCUMENT", "ORDER", "COMMANDE", "N° CDE", "SALES ORDER"], "Client": ["SOLD-TO PARTY", "CUSTOMER", "CLIENT NAME", "NOM CLIENT", "SOLD TO"], "Contrat": ["CONTRACT", "CONTRAT", "PO NUMBER", "PURCHASE ORDER"],
                   "Produit": ["MATERIAL", "PRODUCT", "DESCRIPTION PRODUIT", "MATERIAL DESCRIPTION"], "Quantité commandée": ["ORDER QUANTITY", "QUANTITY", "QUANTITÉ", "QTY"], "Incoterm": ["INCOTERMS", "INCOTERM", "DELIVERY MODE", "MODE DE VENTE"],
                   "Prix unitaire": ["NET PRICE", "UNIT PRICE", "PRIX"], "Devise": ["CURRENCY", "DEVISE", "DOC. CURRENCY"], "Date commande": ["DOCUMENT DATE", "CREATED ON", "DATE COMMANDE"]},
       "deliveries": {"N° livraison": ["DELIVERY", "LIVRAISON", "DELIVERY NO", "OUTBOUND DELIVERY"], "N° commande": ["SALES DOCUMENT", "ORDER", "COMMANDE", "SALES ORDER", "REFERENCE DOCUMENT"], "Navire": ["VESSEL", "SHIP", "NAVIRE"],
                      "Date BL": ["B/L DATE", "BL DATE", "ACTUAL GI DATE", "GOODS ISSUE DATE", "DATE BL"], "Quantité livrée": ["DELIVERY QUANTITY", "QUANTITY", "QTY LOADED", "QUANTITÉ CHARGÉE", "QUANTITÉ"], "Port": ["PORT", "DESTINATION", "DISCHARGE PORT", "PORT OF LOADING"]},
       "invoices": {"N° facture": ["BILLING DOCUMENT", "INVOICE", "FACTURE", "NUMÉRO FACTURE"], "N° livraison": ["DELIVERY", "LIVRAISON", "REFERENCE", "REFERENCE DOCUMENT"], "N° commande": ["SALES DOCUMENT", "ORDER", "COMMANDE", "SALES ORDER"],
                    "Date facture": ["BILLING DATE", "INVOICE DATE", "DATE FACTURE", "DATE CRÉATION FACTURE"], "Montant": ["NET VALUE", "AMOUNT", "MONTANT", "MONTANT FACTURÉE", "TOTAL"], "Échéance": ["DUE DATE", "ECHEANCE", "ÉCHÉANCE", "NET DUE DATE"], "Client": ["PAYER", "CUSTOMER", "SOLD-TO PARTY", "CLIENT NAME"], "Devise": ["CURRENCY", "DEVISE"]},
       "payments": {"N° facture": ["REFERENCE", "INVOICE", "BILLING DOCUMENT", "ASSIGNMENT", "N° FACTURE"], "Montant réglé": ["AMOUNT PAID", "CLEARED AMOUNT", "MONTANT RÈGLEMENT", "MONTANT REGLEMENT 1", "PAYMENT"], "Date règlement": ["CLEARING DATE", "DATE REGLEMENT 1", "DATE RÈGLEMENT", "PAYMENT DATE"], "Solde": ["OPEN AMOUNT", "SOLDE", "SOLDE CLIENT", "AMOUNT IN LOCAL CURRENCY"]}}
cfg = load_config(a.config, DEF)
O = read_any(a.orders, aliases=cfg["orders"], required=["N° commande", "Client"]); D = read_any(a.deliveries, aliases=cfg["deliveries"], required=["N° commande", "Navire", "Date BL"]); I = read_any(a.invoices, aliases=cfg["invoices"], required=["N° facture", "Montant"])
P = read_any(a.payments, aliases=cfg["payments"]) if a.payments else None
ARRETE = pd.Timestamp(a.arrete) if a.arrete else pd.Timestamp.today().normalize()
for df, cols in [(O, ["Quantité commandée", "Prix unitaire"]), (D, ["Quantité livrée"]), (I, ["Montant"])]:
    for c in cols:
        if c in df.columns: df[c] = df[c].map(to_num)
for df, cols in [(O, ["Date commande"]), (D, ["Date BL"]), (I, ["Date facture", "Échéance"])]:
    for c in cols:
        if c in df.columns: df[c] = df[c].map(to_date)
for df in (O, D, I): df["N° commande"] = df["N° commande"].astype(str).str.replace(r"\.0$", "", regex=True).str.strip() if "N° commande" in df.columns else ""
I["N° facture"] = I["N° facture"].astype(str).str.replace(r"\.0$", "", regex=True).str.strip()
if "N° livraison" in D.columns: D["N° livraison"] = D["N° livraison"].astype(str).str.replace(r"\.0$", "", regex=True).str.strip()
if "N° livraison" in I.columns: I["N° livraison"] = I["N° livraison"].astype(str).str.replace(r"\.0$", "", regex=True).str.strip()
D["VK"] = D["Navire"].map(norm_text)
# ---- règlements par facture
if P is not None:
    P["N° facture"] = P["N° facture"].astype(str).str.replace(r"\.0$", "", regex=True).str.strip()
    for c in ["Montant réglé", "Solde"]:
        if c in P.columns: P[c] = P[c].map(to_num)
    if "Date règlement" in P.columns: P["Date règlement"] = P["Date règlement"].map(to_date)
    pay = P.groupby("N° facture").agg(**{"Montant réglé": ("Montant réglé", "sum")} if "Montant réglé" in P.columns else {}, **({"Solde": ("Solde", "sum")} if "Solde" in P.columns else {}), **({"Date règlement": ("Date règlement", "max")} if "Date règlement" in P.columns else {})).reset_index()
    I = I.merge(pay, on="N° facture", how="left")
if "Solde" not in I.columns: I["Solde"] = I["Montant"] - I.get("Montant réglé", pd.Series(0, index=I.index)).fillna(0) if "Montant réglé" in I.columns else np.nan
if "Montant réglé" not in I.columns: I["Montant réglé"] = I["Montant"] - I["Solde"]
# ---- factures → livraison (par n° livraison sinon par n° commande)
inv_by_del = I.groupby("N° livraison").agg(**{"N° factures": ("N° facture", lambda s: " ; ".join(s)), "Montant facturé": ("Montant", "sum"), "Montant réglé": ("Montant réglé", "sum"), "Solde": ("Solde", "sum"), "Date facture": ("Date facture", "min"), "Échéance": ("Échéance", "min"), **({"Date règlement": ("Date règlement", "max")} if "Date règlement" in I.columns else {})}).reset_index() if "N° livraison" in I.columns else pd.DataFrame(columns=["N° livraison"])
inv_by_ord = I.groupby("N° commande").agg(**{"N° factures": ("N° facture", lambda s: " ; ".join(s)), "Montant facturé": ("Montant", "sum"), "Montant réglé": ("Montant réglé", "sum"), "Solde": ("Solde", "sum"), "Date facture": ("Date facture", "min"), "Échéance": ("Échéance", "min"), **({"Date règlement": ("Date règlement", "max")} if "Date règlement" in I.columns else {})}).reset_index() if "N° commande" in I.columns else pd.DataFrame(columns=["N° commande"])
C = D.merge(O.drop_duplicates("N° commande"), on="N° commande", how="left", suffixes=("", " (commande)"))
if "N° livraison" in C.columns and len(inv_by_del): C = C.merge(inv_by_del, on="N° livraison", how="left")
if "Montant facturé" not in C.columns or C["Montant facturé"].isna().all():
    C = C.drop(columns=[c for c in ["N° factures", "Montant facturé", "Montant réglé", "Solde", "Date facture", "Échéance", "Date règlement"] if c in C.columns]).merge(inv_by_ord, on="N° commande", how="left")
else:
    miss = C["Montant facturé"].isna()
    if miss.any() and len(inv_by_ord):
        fill = C.loc[miss, ["N° commande"]].merge(inv_by_ord, on="N° commande", how="left")
        for c in inv_by_ord.columns:
            if c != "N° commande" and c in C.columns: C.loc[miss, c] = fill[c].values
# ---- indicateurs par cargaison
C["Montant attendu"] = (C["Quantité livrée"] * C["Prix unitaire"]) if "Prix unitaire" in C.columns and "Quantité livrée" in C.columns else np.nan
C["Écart quantité"] = (C["Quantité livrée"] - C["Quantité commandée"]) if "Quantité commandée" in C.columns and "Quantité livrée" in C.columns else np.nan
C["Écart montant facturé vs attendu"] = C["Montant facturé"] - C["Montant attendu"]
C["Délai BL → facture (j)"] = (C["Date facture"] - C["Date BL"]).dt.days
C["Délai facture → règlement (j)"] = (C["Date règlement"] - C["Date facture"]).dt.days if "Date règlement" in C.columns else np.nan
C["Retard vs échéance (j)"] = np.where(C["Solde"].fillna(0) > 0.01, (ARRETE - C["Échéance"]).dt.days, np.nan)
C["Tranche retard"] = C["Retard vs échéance (j)"].map(bucket_days)


def statut(r):
    if pd.isna(r["Montant facturé"]) or r["Montant facturé"] == 0: return "Livré non facturé"
    if pd.isna(r["Solde"]) or r["Solde"] <= 0.01: return "Facturé et encaissé"
    if r["Montant réglé"] and r["Montant réglé"] > 0.01: return "Facturé partiellement encaissé"
    return "Facturé non encaissé"


C["Statut O2C"] = C.apply(statut, axis=1)
C["Alerte"] = np.where(C["Statut O2C"] == "Livré non facturé", np.where((ARRETE - C["Date BL"]).dt.days > a.delai_facture, f"Non facturé > {a.delai_facture} j après BL", "À facturer"),
                np.where(C["Retard vs échéance (j)"].fillna(0) > 90, "Retard > 90 j", np.where(C["Écart montant facturé vs attendu"].abs() > np.maximum(1, 0.005 * C["Montant attendu"].abs().fillna(0)), "Montant facturé ≠ quantité × prix", "")))
C["Mois BL"] = C["Date BL"].dt.to_period("M").astype(str)
cols = [c for c in ["Navire", "Date BL", "Mois BL", "N° livraison", "N° commande", "Contrat", "Client", "Produit", "Incoterm", "Port", "Quantité commandée", "Quantité livrée", "Écart quantité", "Prix unitaire", "Devise", "Montant attendu", "N° factures", "Date facture", "Montant facturé", "Écart montant facturé vs attendu", "Échéance", "Montant réglé", "Date règlement", "Solde", "Délai BL → facture (j)", "Délai facture → règlement (j)", "Retard vs échéance (j)", "Tranche retard", "Statut O2C", "Alerte"] if c in C.columns]
C = C[cols].sort_values(["Date BL", "Navire"]).reset_index(drop=True)
# ---- classeur
wb = Workbook(); ws = wb.active; ws.title = "Synthèse"
ws["A1"] = "Order-to-cash par cargaison — commande → BL → facture → encaissement"; ws["A1"].font = TITLE
ws["A2"] = f"Arrêté : {ARRETE:%d/%m/%Y}. Une ligne par livraison (navire + BL). Formules sur l'onglet « Cargaisons »."; ws["A2"].font = NOTE
cm = colmap(C); n = len(C); R_ = lambda c: rng("Cargaisons", cm[c], n)
r = kpi_rows(ws, 4, [("Cargaisons", f'=COUNTA({R_("Navire")})', "0"), ("Quantité livrée", f'=SUM({R_("Quantité livrée")})' if "Quantité livrée" in cm else 0, MONEY0),
                    ("Montant facturé", f'=SUM({R_("Montant facturé")})', MONEY0), ("Montant réglé", f'=SUM({R_("Montant réglé")})', MONEY0), ("Solde client ouvert", f'=SUM({R_("Solde")})', MONEY0),
                    ("Montant attendu (qté × prix)", f'=SUM({R_("Montant attendu")})' if "Montant attendu" in cm else 0, MONEY0), ("Écart facturé vs attendu", f'=SUM({R_("Écart montant facturé vs attendu")})' if "Écart montant facturé vs attendu" in cm else 0, MONEY0),
                    ("Délai médian BL → facture (j)", f'=MEDIAN({R_("Délai BL → facture (j)")})', "0"), ("Délai médian facture → règlement (j)", f'=MEDIAN({R_("Délai facture → règlement (j)")})' if "Délai facture → règlement (j)" in cm else 0, "0")], "1. Volumes et montants")
r = kpi_rows(ws, r, [(s, f'=COUNTIF({R_("Statut O2C")},"{s}")', "0") for s in ["Livré non facturé", "Facturé non encaissé", "Facturé partiellement encaissé", "Facturé et encaissé"]] + [("Solde des cargaisons " + s.lower(), f'=SUMIFS({R_("Solde")},{R_("Statut O2C")},"{s}")', MONEY0) for s in ["Facturé non encaissé", "Facturé partiellement encaissé"]], "2. Statut O2C (cargaisons)")
r = kpi_rows(ws, r, [(b, f'=SUMIFS({R_("Solde")},{R_("Tranche retard")},"{b}")', MONEY0) for b in BUCKETS], "3. Encours vieilli (solde par tranche de retard vs échéance)")
r = kpi_rows(ws, r, [(al, f'=COUNTIF({R_("Alerte")},"{al}")', "0") for al in [f"Non facturé > {a.delai_facture} j après BL", "À facturer", "Retard > 90 j", "Montant facturé ≠ quantité × prix"]], "4. Alertes")
ws.column_dimensions["A"].width = 48; ws.column_dimensions["B"].width = 18; ws.sheet_view.showGridLines = False
ws2 = wb.create_sheet("Cargaisons")
write_df(ws2, C, money=[c for c in C.columns if "Montant" in c or c in ("Solde", "Prix unitaire")], date=["Date BL", "Date facture", "Échéance", "Date règlement"], intc=["Délai BL → facture (j)", "Délai facture → règlement (j)", "Retard vs échéance (j)"], widths={"Alerte": 34, "Statut O2C": 28, "N° factures": 28})
last = get_column_letter(len(C.columns)); sc = cm["Statut O2C"]
for s_, f in [("Livré non facturé", FILL_RED), ("Facturé non encaissé", FILL_YEL), ("Facturé et encaissé", FILL_GRN)]: ws2.conditional_formatting.add(f"A2:{last}{n + 1}", FormulaRule(formula=[f'${sc}2="{s_}"'], fill=f))
mois = C.groupby("Mois BL").agg(Cargaisons=("Navire", "size"), Facturé=("Montant facturé", "sum"), Réglé=("Montant réglé", "sum"), Solde=("Solde", "sum"), Délai_BL_facture=("Délai BL → facture (j)", "median")).reset_index()
write_df(wb.create_sheet("Par mois"), mois, money0=["Facturé", "Réglé", "Solde"], intc=["Cargaisons", "Délai_BL_facture"], autofilter=False)
cli = C.groupby("Client").agg(Cargaisons=("Navire", "size"), Facturé=("Montant facturé", "sum"), Solde=("Solde", "sum"), Non_facturé=("Statut O2C", lambda s: (s == "Livré non facturé").sum())).sort_values("Solde", ascending=False).reset_index()
write_df(wb.create_sheet("Par client"), cli, money0=["Facturé", "Solde"], intc=["Cargaisons", "Non_facturé"], autofilter=False)
summary = {"module": "o2c", "arrete": str(ARRETE.date()), "cargaisons": n, "facture": float(C["Montant facturé"].sum()), "regle": float(C["Montant réglé"].sum()), "solde": float(C["Solde"].sum()), "non_factures": int((C["Statut O2C"] == "Livré non facturé").sum()),
           "statuts": C["Statut O2C"].value_counts().to_dict(), "alertes": C["Alerte"].value_counts().to_dict(), "delai_bl_facture_median": float(C["Délai BL → facture (j)"].median()) if C["Délai BL → facture (j)"].notna().any() else None,
           "delai_facture_reglement_median": float(C["Délai facture → règlement (j)"].median()) if "Délai facture → règlement (j)" in C.columns and C["Délai facture → règlement (j)"].notna().any() else None, "par_mois": mois.to_dict("records"), "top_clients_solde": cli.head(5).to_dict("records")}
finish(wb, a.out, summary)
print(json.dumps({k: summary[k] for k in ["cargaisons", "facture", "regle", "solde", "non_factures", "statuts"]}, ensure_ascii=False, default=str))
