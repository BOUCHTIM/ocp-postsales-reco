"""Approximate pptx -> HTML renderer for visual QA (no LibreOffice available).
Geometry is exact (EMU -> px at 96 dpi); fonts fall back to system (Calibri/Cambria -> Carlito/Caladea or Helvetica)."""
import sys, html
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR

THEME = {"dk1": "0B1F3A", "lt1": "FFFFFF", "dk2": "4A5568", "lt2": "EEF3F8", "accent1": "0F4C81", "accent2": "2A9D8F",
         "accent3": "C8553D", "accent4": "E9C46A", "accent5": "7FB7D6", "accent6": "A0AEC0"}
TC = {MSO_THEME_COLOR.TEXT_1: "dk1", MSO_THEME_COLOR.TEXT_2: "dk2", MSO_THEME_COLOR.BACKGROUND_1: "lt1", MSO_THEME_COLOR.BACKGROUND_2: "lt2",
      MSO_THEME_COLOR.ACCENT_1: "accent1", MSO_THEME_COLOR.ACCENT_2: "accent2", MSO_THEME_COLOR.ACCENT_3: "accent3",
      MSO_THEME_COLOR.ACCENT_4: "accent4", MSO_THEME_COLOR.ACCENT_5: "accent5", MSO_THEME_COLOR.ACCENT_6: "accent6",
      MSO_THEME_COLOR.DARK_1: "dk1", MSO_THEME_COLOR.LIGHT_1: "lt1", MSO_THEME_COLOR.DARK_2: "dk2", MSO_THEME_COLOR.LIGHT_2: "lt2"}
PX = 96 / 914400.0


def color_of(fmt, default=None):
    try:
        if fmt.type is None:
            return default
        if fmt.type == 1:  # RGB
            return str(fmt.rgb)
        if fmt.type == 2:  # scheme
            return THEME.get(TC.get(fmt.theme_color, ""), default)
    except Exception:
        return default
    return default


def fill_color(shape):
    try:
        f = shape.fill
        if f.type == 1:
            return color_of(f.fore_color, None)
    except Exception:
        pass
    return None


def line_css(shape):
    try:
        ln = shape.line
        if ln.fill.type == 1:
            c = color_of(ln.color, "000000")
            w = (ln.width or Emu(9525)) * PX
            return f"border:{max(1, w):.1f}px solid #{c};"
    except Exception:
        pass
    return ""


def text_html(tf, default_color, default_size=18.0, default_font="Calibri"):
    out = []
    for p in tf.paragraphs:
        align = {1: "left", 2: "center", 3: "right"}.get(p.alignment, "left") if p.alignment else "left"
        runs = []
        for r in p.runs:
            f = r.font
            size = f.size.pt if f.size else default_size
            col = color_of(f.color, default_color) if f.color is not None else default_color
            b = "font-weight:bold;" if f.bold else ""
            i = "font-style:italic;" if f.italic else ""
            fn = f.name or default_font
            runs.append(f'<span style="font-size:{size * 96 / 72:.1f}px;color:#{col};{b}{i}font-family:\'{fn}\',Carlito,Caladea,Helvetica,Arial,sans-serif">{html.escape(r.text)}</span>')
        bullet = ""
        try:
            if p._p.pPr is not None and p._p.pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None:
                bullet = "• "
        except Exception:
            pass
        sa = ""
        try:
            if p.space_after is not None:
                sa = f"margin-bottom:{p.space_after.pt * 96 / 72:.1f}px;"
        except Exception:
            pass
        out.append(f'<div style="text-align:{align};{sa}white-space:pre-wrap;line-height:1.2">{bullet}{"".join(runs) or "&nbsp;"}</div>')
    return "".join(out)


def shape_html(sh, inherit=None):
    try:
        x, y, w, h = sh.left, sh.top, sh.width, sh.height
    except Exception:
        x = y = w = h = None
    if (x is None or w is None) and inherit is not None:
        x, y, w, h = inherit.left, inherit.top, inherit.width, inherit.height
    if x is None:
        return ""
    style = f"position:absolute;left:{x * PX:.1f}px;top:{y * PX:.1f}px;width:{w * PX:.1f}px;height:{h * PX:.1f}px;box-sizing:border-box;"
    if sh.shape_type == MSO_SHAPE_TYPE.TABLE:
        rows = []
        for r in sh.table.rows:
            cells = []
            for c in r.cells:
                fc = None
                try:
                    if c.fill.type == 1:
                        fc = color_of(c.fill.fore_color)
                except Exception:
                    pass
                cells.append(f'<td style="border:1px solid #{THEME["lt2"]};padding:4px;vertical-align:middle;{"background:#" + fc + ";" if fc else ""}">{text_html(c.text_frame, THEME["dk1"], 11)}</td>')
            rows.append("<tr>" + "".join(cells) + "</tr>")
        cols = "".join(f'<col style="width:{c.width * PX:.1f}px">' for c in sh.table.columns)
        return f'<table style="{style}border-collapse:collapse;table-layout:fixed;outline:1px dashed #f0f">{cols}{"".join(rows)}</table>'
    if sh.shape_type == MSO_SHAPE_TYPE.CHART:
        return f'<div style="{style}background:repeating-linear-gradient(45deg,#f6f6f6,#f6f6f6 10px,#eee 10px,#eee 20px);border:1px dashed #999;display:flex;align-items:center;justify-content:center;color:#666;font:12px Helvetica">CHART ({sh.chart.chart_type})</div>'
    fc = fill_color(sh)
    bg = f"background:#{fc};" if fc else ""
    radius = ""
    try:
        if sh.auto_shape_type is not None and "ROUNDED" in str(sh.auto_shape_type):
            radius = "border-radius:8px;"
        if sh.auto_shape_type is not None and "OVAL" in str(sh.auto_shape_type):
            radius = "border-radius:50%;"
        if sh.auto_shape_type is not None and "ARROW" in str(sh.auto_shape_type):
            radius = "clip-path:polygon(0 25%,60% 25%,60% 0,100% 50%,60% 100%,60% 75%,0 75%);"
    except Exception:
        pass
    if sh.shape_type == MSO_SHAPE_TYPE.LINE or (h == 0 and w):
        return f'<div style="{style}height:1px;background:#{THEME["lt2"]}"></div>'
    inner = ""
    if sh.has_text_frame and sh.text_frame.text.strip():
        tf = sh.text_frame
        bp = tf.margin_left if tf.margin_left is not None else Emu(91440)
        tpad = tf.margin_top if tf.margin_top is not None else Emu(45720)
        va = {1: "flex-start", 2: "center", 3: "flex-end"}.get(tf.vertical_anchor, "flex-start") if tf.vertical_anchor else "flex-start"
        pad = f"padding:{tpad * PX:.1f}px {bp * PX:.1f}px;"
        dc = THEME["dk1"]
        inner = f'<div style="position:absolute;inset:0;{pad}display:flex;flex-direction:column;justify-content:{va};overflow:visible;outline:1px dashed rgba(255,0,255,.35)">{text_html(tf, dc)}</div>'
    return f'<div style="{style}{bg}{radius}{line_css(sh)}">{inner}</div>'


def render(path, out_prefix):
    prs = Presentation(path)
    W, Hh = prs.slide_width * PX, prs.slide_height * PX
    pages = []
    for i, slide in enumerate(prs.slides, 1):
        lay = slide.slide_layout
        bg = "FFFFFF"
        try:
            if lay.background.fill.type == 1:
                bg = color_of(lay.background.fill.fore_color, "FFFFFF")
        except Exception:
            pass
        parts = []
        for sh in lay.shapes:
            if sh.is_placeholder:
                continue
            parts.append(shape_html(sh))
        for sh in slide.shapes:
            inh = None
            if sh.is_placeholder:
                try:
                    inh = lay.placeholders.get(idx=sh.placeholder_format.idx)
                except Exception:
                    inh = None
            parts.append(shape_html(sh, inh))
        page = f'<div style="position:relative;width:{W:.0f}px;height:{Hh:.0f}px;background:#{bg};overflow:hidden;margin:16px auto;box-shadow:0 0 0 1px #999">{"".join(parts)}</div>'
        pages.append(page)
        with open(f"{out_prefix}-{i:02d}.html", "w") as f:
            f.write(f'<!doctype html><meta charset="utf-8"><body style="margin:0;background:#ddd">{page}</body>')
    with open(f"{out_prefix}-all.html", "w") as f:
        f.write(f'<!doctype html><meta charset="utf-8"><body style="margin:0;background:#ddd">{"".join(pages)}</body>')
    print("pages", len(pages))


if __name__ == "__main__":
    render(sys.argv[1], sys.argv[2])
