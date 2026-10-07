"""Briques communes aux modules : lecture tolérante (alias de colonnes), normalisation, écriture Excel homogène (Arial, en-têtes, formats),
synthèse en formules, journal JSON. Importé par o2c.py, lc.py, aging.py, claims_register.py, laytime.py, dashboard.py."""
import re, json, os
import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864"); HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BASE = Font(name=FONT, size=10); BOLD = Font(name=FONT, size=10, bold=True); TITLE = Font(name=FONT, size=14, bold=True, color="1F3864"); NOTE = Font(name=FONT, size=9, italic=True)
THIN = Side(style="thin", color="BFBFBF"); BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
MONEY = '#,##0.00;[Red](#,##0.00);"-"'; MONEY0 = '#,##0;[Red](#,##0);"-"'; PCT = "0.0%"; DATE = "DD/MM/YYYY"
FILL_RED = PatternFill("solid", fgColor="F8CBAD"); FILL_YEL = PatternFill("solid", fgColor="FFE699"); FILL_GRN = PatternFill("solid", fgColor="E2EFDA"); FILL_GREY = PatternFill("solid", fgColor="EEF3F8")


def norm_name(c): return re.sub(r"\s+", " ", str(c)).strip().upper()


def norm_text(v):
    if pd.isna(v): return ""
    s = str(v).upper(); s = re.sub(r"\bM/?V\b", " ", s); s = re.sub(r"[^A-Z0-9 ]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def to_date(v):
    if pd.isna(v) or v == "": return pd.NaT
    if isinstance(v, (pd.Timestamp,)): return v.normalize()
    s = str(v); m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        try: return pd.Timestamp(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except Exception: return pd.NaT
    return pd.to_datetime(s, errors="coerce").normalize() if pd.notna(pd.to_datetime(s, errors="coerce")) else pd.NaT


def to_num(v):
    if pd.isna(v): return np.nan
    if isinstance(v, (int, float, np.integer, np.floating)): return float(v)
    s = str(v).replace(" ", "").replace(" ", "").replace(",", ".")
    try: return float(s)
    except ValueError: return np.nan


def read_any(path, sheet=None, aliases=None, required=(), nrows=None):
    """Lit un xlsx/csv ; 'fichier.xlsx:Onglet' choisit l'onglet (sous-chaîne acceptée) ; renomme les colonnes via aliases {canon: [alias...]} ; vérifie les colonnes requises."""
    path = str(path)
    if not os.path.exists(path) and ":" in path:
        path, _, sh = path.rpartition(":"); sheet = sheet or sh
    if path.lower().endswith(".csv"): df = pd.read_csv(path, sep=None, engine="python", nrows=nrows)
    else:
        xl = pd.ExcelFile(path)
        if sheet is None or sheet not in xl.sheet_names:
            sheet = next((s for s in xl.sheet_names if sheet and norm_name(sheet) in norm_name(s)), xl.sheet_names[0])
        df = pd.read_excel(path, sheet_name=sheet, nrows=nrows)
    df = rename_cols(df, aliases or {})
    missing = [c for c in required if c not in df.columns]
    if missing: raise SystemExit(f"{os.path.basename(path)} : colonnes manquantes {missing}. Colonnes présentes : {list(df.columns)}. Ajouter des alias (--config) ou utiliser le modèle fourni.")
    for c in df.columns:
        if pd.api.types.is_string_dtype(df[c]) or df[c].dtype == object:
            df[c] = df[c].map(lambda v: re.sub(r"_x[0-9A-Fa-f]{4}_|[\x00-\x08\x0b-\x1f]", "", v) if isinstance(v, str) else v)
    return df


def rename_cols(df, mapping):
    cols = {norm_name(c): c for c in df.columns}; ren = {}
    for canon, al in mapping.items():
        if canon in df.columns: continue
        for x in [canon] + list(al):
            if norm_name(x) in cols: ren[cols[norm_name(x)]] = canon; break
    return df.rename(columns=ren)


def load_config(path, defaults):
    cfg = json.loads(json.dumps(defaults))
    if path:
        user = json.load(open(path))
        for k, v in user.items():
            if isinstance(v, dict) and isinstance(cfg.get(k), dict): cfg[k].update(v)
            else: cfg[k] = v
    return cfg


def write_df(ws, df, r0=1, money=(), money0=(), date=(), pct=(), intc=(), widths=None, autofilter=True):
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(row=r0, column=j, value=str(c)); cell.font = HDR_FONT; cell.fill = HDR_FILL; cell.border = BORDER; cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, rec in enumerate(df.itertuples(index=False), r0 + 1):
        for j, v in enumerate(rec, 1):
            if isinstance(v, float) and np.isnan(v): v = None
            if isinstance(v, pd.Timestamp): v = None if pd.isna(v) else v.to_pydatetime()
            if v is pd.NaT: v = None
            if hasattr(v, "item") and not isinstance(v, (str, bytes)):
                try: v = v.item()
                except Exception: pass
            cell = ws.cell(row=i, column=j, value=v); cell.font = BASE; cell.border = BORDER
            col = df.columns[j - 1]
            if col in money: cell.number_format = MONEY
            elif col in money0: cell.number_format = MONEY0
            elif col in date: cell.number_format = DATE
            elif col in pct: cell.number_format = PCT
            elif col in intc: cell.number_format = "0"
    ws.row_dimensions[r0].height = 40
    if autofilter and len(df): ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(df.columns))}{r0 + len(df)}"
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    for j, c in enumerate(df.columns, 1):
        w = (widths or {}).get(c)
        if w is None:
            mx = max([len(str(c)) // 2 + 4] + [len(str(x)) for x in df[c].head(300).tolist() if x is not None and not (isinstance(x, float) and np.isnan(x))])
            w = min(max(9, mx + 2), 45)
        ws.column_dimensions[get_column_letter(j)].width = w
    return r0 + len(df) + 2


def kpi_rows(ws, r, rows, title=None):
    """rows = [(label, formule_ou_valeur, format)] ; écrit un bloc d'indicateurs."""
    if title: ws.cell(row=r, column=1, value=title).font = BOLD; r += 1
    for lab, f, fmt in rows:
        ws.cell(row=r, column=1, value=lab).font = BASE
        c = ws.cell(row=r, column=2, value=f); c.font = BASE; c.border = BORDER; c.number_format = fmt
        r += 1
    return r + 1


def colmap(df): return {c: get_column_letter(i + 1) for i, c in enumerate(df.columns)}


def rng(sheet, col_letter, n): return f"'{sheet}'!${col_letter}$2:${col_letter}${n + 1}"


def bucket_days(d, edges=(0, 30, 60, 90, 180)):
    if pd.isna(d): return ""
    if d <= edges[0]: return "Non échu"
    for lo, hi in zip(edges, edges[1:]):
        if d <= hi: return f"{lo + 1}-{hi} j"
    return f"> {edges[-1]} j"


BUCKETS = ["Non échu", "1-30 j", "31-60 j", "61-90 j", "91-180 j", "> 180 j"]


def finish(wb, out, summary=None):
    wb.calculation.fullCalcOnLoad = True; wb.save(out)
    if summary is not None:
        def clean(o):
            if isinstance(o, dict): return {str(k): clean(v) for k, v in o.items()}
            if isinstance(o, list): return [clean(v) for v in o]
            if isinstance(o, float) and o != o: return None
            if isinstance(o, (pd.Timestamp,)): return None if pd.isna(o) else str(o.date())
            if hasattr(o, "item"):
                try: return clean(o.item())
                except Exception: return str(o)
            return o
        json.dump(clean(summary), open(os.path.splitext(out)[0] + "_summary.json", "w"), ensure_ascii=False, indent=1, default=str, allow_nan=False)
    print("saved", os.path.basename(out))
